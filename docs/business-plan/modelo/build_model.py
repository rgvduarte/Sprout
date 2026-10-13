#!/usr/bin/env python3
"""Modelo financeiro do Sprout em Excel, com fórmulas vivas.

Lê `pressupostos.json` (mesma pasta) e escreve `../Sprout-Modelo-Financeiro.xlsx`:
todas as saídas são fórmulas que dependem das células de input da folha
"Pressupostos"; nenhum resultado é escrito como número.

Depois recalcula o livro com o LibreOffice em modo headless, confirma que não há
erros de fórmula, lê os valores recalculados e grava `resultados.json` (a fonte
única de números para os PDF). Também verifica o livro contra
`calculo_referencia.py` (três cenários do documento-fonte) e contra um espelho
em Python do mesmo motor (quatro cenários e sensibilidades).

Uso:  python3 build_model.py            # constrói, recalcula, verifica, grava JSON
      python3 build_model.py --no-recalc  # só constrói o .xlsx

Arquitetura (uma só definição do cálculo, duas disposições):
- as linhas do motor (volumes, economia por kg, receitas, custos, pessoal, DR,
  cash flow, capital) são definidas uma vez em `define_lines()`;
- disposição "detalhe": o cenário ativo (seletor em Pressupostos!C4) espalhado
  pelas folhas Volumes, Receitas, Pessoal, Custos, DR, Cash flow e Capital;
- disposição "Motor": as mesmas linhas em blocos fixos, um por cenário e um por
  caso de sensibilidade; Cenários, Unit economics e Sensibilidade leem daqui.
"""
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
JSON_PATH = HERE / "pressupostos.json"
OUT_XLSX = HERE.parent / "Sprout-Modelo-Financeiro.xlsx"
OUT_JSON = HERE / "resultados.json"

SCEN = ["pessimista", "base", "otimista", "recomendado"]
SCEN_LABEL = ["Pessimista", "Base", "Otimista", "Recomendado"]
SCEN_LONG = ["Pessimista", "Base", "Otimista", "Recomendado (reestruturado)"]
YEARS = [2027, 2028, 2029, 2030, 2031]
YC = ["C", "D", "E", "F", "G"]          # colunas dos anos em todas as folhas de cálculo
NY = 5

# ----------------------------------------------------------------------------- estilo
ARIAL = "Arial"
INK, MUTED, RULE, PANEL, ACCENT = "1A1E16", "5B6352", "D7DBCC", "F4F5EF", "3F5A1E"
BLUE, GREEN, BLACK = "0000FF", "008000", "000000"
WARN, CRIT, OK = "A86A12", "B42318", "2E7D4F"

FMT_EUR = '#,##0;(#,##0);"-"'
FMT_EKG = '0.000;(0.000);"-"'
FMT_EKG2 = '0.00;(0.00);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_PCT2 = '0.00%;(0.00%);"-"'
FMT_T = '#,##0.0;(#,##0.0);"-"'
FMT_INT = '#,##0;(#,##0);"-"'
FMT_DAYS = '0;(0);"-"'
FMT_FTE = '0.0;(0.0);"-"'
FMT_NUM = '0.00'
FMT_YEAR = '0'
FMT_DATE = 'mmm yyyy'
FMT_MES = '0.0'

THIN = Side(style="thin", color=RULE)
B_BOTTOM = Border(bottom=THIN)
B_TOP = Border(top=THIN)
FILL_PANEL = PatternFill("solid", fgColor=PANEL)
FILL_SEL = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="FAFBF7")


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=ARIAL, size=size, color=color, bold=bold, italic=italic)


def q(sheet):
    """Nome de folha pronto para referência entre folhas."""
    if any(ch in sheet for ch in " &-'()"):
        return "'" + sheet.replace("'", "''") + "'"
    return sheet


def is_link(formula):
    """Fórmula que é só uma ligação a outra folha (fica a verde)."""
    import re
    return bool(re.fullmatch(r"=('[^']+'|[A-Za-z][\w]*)!\$?[A-Z]{1,3}\$?\d+", formula))


def put(ws, coord, value, fmt=None, bold=False, kind=None, italic=False, color=None, align=None, fill=None):
    """Escreve uma célula com a convenção de cores: input azul, fórmula preta, ligação verde."""
    c = ws[coord]
    c.value = value
    if kind == "input":
        col = BLUE
    elif isinstance(value, str) and value.startswith("="):
        col = GREEN if is_link(value) else BLACK
    else:
        col = color or INK
    if color and kind != "input":
        col = color
    c.font = font(col, bold=bold, italic=italic)
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if fill:
        c.fill = fill
    return c


def title(ws, text, sub=None):
    ws["A1"].value = text
    ws["A1"].font = Font(name=ARIAL, size=14, bold=True, color=INK)
    if sub:
        ws["A2"].value = sub
        ws["A2"].font = font(MUTED, italic=True)
    ws.sheet_view.showGridLines = False


def header_row(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = font(INK, bold=True)
        c.fill = FILL_PANEL
        c.border = B_BOTTOM
        c.alignment = Alignment(horizontal="left" if i < 2 else "right", vertical="center", wrap_text=True)


def section_row(ws, row, text, ncols=8):
    c = ws.cell(row=row, column=1, value=text)
    c.font = font(ACCENT, bold=True)
    for j in range(1, ncols + 1):
        ws.cell(row=row, column=j).border = B_BOTTOM


WRAP = Alignment(wrap_text=True, vertical="top")


# ----------------------------------------------------------------------------- pressupostos
class Param:
    def __init__(self, pid, sec, label, unit, kind, vals, fmt, fonte):
        assert kind in ("const", "scen", "cseries", "sseries"), pid
        self.pid, self.sec, self.label, self.unit = pid, sec, label, unit
        self.kind, self.vals, self.fmt, self.fonte = kind, vals, fmt, fonte
        self.row = None          # const/scen/cseries
        self.rows = None         # sseries: [r_pess, r_base, r_otim, r_recom, r_ativo]


def scen_vals(d):
    return {s: d[s] for s in SCEN}


def build_params(P):
    az = P["commodities"]["azeite"]
    H, D, G = (az["canais"][k] for k in ("horeca_gourmet", "distribuidores_export", "granel_certificado"))
    cfc = az["custos_fisicos_e_custodia"]
    ag, cf = P["commodities"]["agua"], P["commodities"]["cafe"]
    ft, dd = P["receitas"]["financiamento_stock_terceiros"], P["receitas"]["dados_e_indice_eur"]
    fin = P["financiamento"]
    spv = fin["custos_spv"]
    cu = P["custos"]
    pe, ox, sc, kyc = cu["pessoal"], cu["opex_nao_salarial_eur"], cu["seguro_credito"], cu["kyc"]
    cx, imp, kp, st = P["capex"], P["impostos"], P["kpis_piloto"], P["stress_tests"]
    ue, sens, tese = P["unit_economics"], P["sensibilidades"], az["tese_documento_fonte"]

    L = []

    def p(*a):
        L.append(Param(*a))

    S1 = "Azeite: preço do granel e prémios de origem"
    p("az_indice", S1, "Índice do granel AOVE (média 2 semanas SIMA + Poolred)", "€/kg", "scen", scen_vals(az["preco_indice_granel_eur_kg"]), FMT_EKG2,
      "[F] SIMA 3,40–3,80 e Poolred 3,34 (out. 2026); 3,50 arredondado para cima; constante (compra back-to-back). 00 §1.2")
    p("h_premio_lagar", S1, "Prémio de origem pago ao lagar — HoReCa/gourmet", "€/kg", "scen", scen_vals(H["premio_lagar_eur_kg"]), FMT_EKG2,
      "[E] lote individualizado com ficha analítica; Recomendado 0,15 (DOP/monovarietal, o prémio mais alto do conjunto verificado)")
    p("d_premio_lagar", S1, "Prémio de origem pago ao lagar — distribuidores/exportação", "€/kg", "scen", scen_vals(D["premio_lagar_eur_kg"]), FMT_EKG2,
      "[E] igual ao HoReCa nos cenários do documento-fonte (calculo_referencia.py usa um só prémio)")
    p("g_premio_lagar", S1, "Prémio de origem pago ao lagar — granel certificado", "€/kg", "scen", scen_vals(G["premio_lagar_eur_kg"]), FMT_EKG2,
      "[E] granel comprado ao índice: o lagar ganha prazo (≤ 30 dias por SCT Inst), não preço")

    S2 = "Azeite: canal HoReCa/gourmet"
    p("h_ref", S2, "Referência do comprador (cotação documentada do equivalente, s/ IVA)", "€/kg", "scen", scen_vals(H["preco_referencia_eur_kg"]), FMT_EKG2,
      "[E, NV] equivalente de retalho standard 5,3–6,0; 6,00 só com origem identificada documentada (gate: ≥ 20 cotações). 00 §1.4; 01 §2.1")
    p("h_desc", S2, "Desconto ao comprador sobre a referência", "%", "scen", scen_vals(H["desconto_efetivo"]), FMT_PCT,
      "Documento-fonte 10–15%; 10% porque 15% leva a margem para ~1,4%. 01 §2.5")
    p("h_dr", S2, "Prazo médio de recebimento", "dias", "scen", scen_vals(H["dias_recebimento"]), FMT_DAYS,
      "[F] restauração: 30 dias de atraso médio, 11% paga a horas (Informa D&B mai. 2026); 35 dias com SDD B2B")
    p("h_prov", S2, "Provisão de crédito", "% vendas", "scen", scen_vals(H["provisao_credito_pct_vendas"]), FMT_PCT2,
      "[E] SDD B2B sem direito a reembolso + limites por cliente")
    S3 = "Azeite: canal distribuidores/exportação"
    p("d_ref", S3, "Referência = preço de saída do embalador", "€/kg", "scen", scen_vals(D["preco_referencia_eur_kg"]), FMT_EKG2,
      "[E, NV] granel 3,50 + embalamento 0,47–0,51 + margem do embalador 0,19–0,50 = 4,16–4,51. 00 §1.4")
    p("d_premio_venda", S3, "Prémio de origem verificada pago pelo comprador", "€/kg", "scen", scen_vals(D["premio_origem_venda_eur_kg"]), FMT_EKG2,
      "[E] 0 no documento-fonte; Recomendado 0,10 = limite inferior da condição para a seed (0,10–0,20). 01 §5.4")
    p("d_desc", S3, "Desconto ao comprador", "%", "scen", scen_vals(D["desconto_efetivo"]), FMT_PCT,
      "0%: qualquer desconto põe o canal abaixo do break-even (4,29 €/kg). 01 §2.5")
    p("d_ds", S3, "Stock médio", "dias", "scen", scen_vals(D["dias_stock"]), FMT_DAYS, "[E] paletes levantadas em ~60 dias")
    p("d_dr", S3, "Prazo médio de recebimento", "dias", "scen", scen_vals(D["dias_recebimento"]), FMT_DAYS,
      "[E] Recomendado 55 dias (mais exportação embalada; valor do pessimista)")
    p("d_prov", S3, "Provisão de crédito", "% vendas", "scen", scen_vals(D["provisao_credito_pct_vendas"]), FMT_PCT2, "[E]")

    S4 = "Outros pressupostos por cenário"
    p("ag_com", S4, "Água: comissão de agente sobre o GMV", "%", "scen", scen_vals(ag["comissao_agente_pct_gmv"]), FMT_PCT,
      "[E, NV] banda 0,5–3%; 4–5% só com cotação real de um engarrafador. 01 §6.1")
    p("cf_inicio", S4, "Café: ano de início", "ano", "scen", scen_vals(cf["ano_inicio"]), FMT_YEAR,
      "[F] EUDR aplicável a 30-12-2026 (grandes/médios) e 30-06-2027 (micro/pequenos); Recomendado adia para 2030")
    p("cf_hedge", S4, "Café: custo de cobertura (forward EUR/USD; ICE > 50–100 t/ano)", "% compras", "scen", scen_vals(cf["custo_cobertura_pct_compras"]), FMT_PCT2,
      "[E/NV] forward 0,1–0,5%; ICE < 0,5% do nocional por semestre. 40 §2.3–2.4")
    p("tp_share", S4, "Verificação para terceiros: parte dos custos suportada pela Sprout", "%", "scen",
      scen_vals(ft["pct_custo_verificacao_suportado_sprout"]), FMT_PCT,
      "1 = Sprout paga collateral manager e análises; 0 = banco/mutuário paga (alternativa da nota de custos)")
    p("fin_util", S4, "Utilização média do capital levantado nas séries", "%", "scen", scen_vals(fin["utilizacao_media_capital"]), FMT_PCT,
      "[E] capital levantado a 12–18 meses vs. rotação do stock de 2–4 meses. 40 §5.1")
    p("fin_pre_seed", S4, "Ronda pré-seed (recebida até 31-01-2027)", "€", "scen", scen_vals(fin["ronda_pre_seed_eur"]), FMT_EUR,
      "01 §5.5 e §7.3: 1,1 M€ (ou 0,8 M€ + bridge firme de 0,3 M€)")

    # ------------------------------------------------------------ constantes
    C1 = "Azeite: canal HoReCa/gourmet (constantes)"
    p("h_emb", C1, "Embalamento no lagar (5 L PET/BIB e vidro)", "€/kg", "const", H["embalamento_eur_kg"], FMT_EKG2,
      "[F] Observatorio 0,47–0,51 €/kg + 0,05 pelo mix com vidro. 00 §1.4")
    p("h_3pl", C1, "Armazém logístico (3PL) e picking", "€/kg", "const", H["armazem_logistico_picking_eur_kg"], FMT_EKG2,
      "[E, NV] encomendas de 20–50 kg a ~2–4 € cada")
    p("h_lote", C1, "Dimensão do lote-tipo", "t", "const", H["lote_medio_t"], FMT_T, "01 §2.2")
    p("h_anal_n", C1, "Análises por lote (entrada + intermédia)", "n.º", "const", H["analises_por_lote"], FMT_INT, "40 §5.2")
    p("h_cust", C1, "Custódia (collateral manager) por lote", "€/lote", "const", H["inspecao_collateral_manager_eur_por_lote"], FMT_EUR,
      "[E, NV] 300–500 € por visita: entrada + ~3 tranches + fecho")
    p("h_pag", C1, "Pagamentos bancários (SEPA)", "€/kg", "const", H["pagamentos_bancarios_eur_kg"], FMT_EKG,
      "[E] comissões por operação; KYC e faturação no custo fixo de tecnologia")
    C2 = "Azeite: canal distribuidores/exportação (constantes)"
    p("d_emb", C2, "Embalamento no lagar (paletes de 5 L)", "€/kg", "const", D["embalamento_eur_kg"], FMT_EKG2, "[F] Observatorio 0,47–0,51 €/kg")
    p("d_3pl", C2, "Armazém logístico (3PL)", "€/kg", "const", D["armazem_logistico_picking_eur_kg"], FMT_EKG2, "Levantamento no lagar")
    p("d_lote", C2, "Dimensão do lote-tipo", "t", "const", D["lote_medio_t"], FMT_T, "01 §2.5")
    p("d_anal_n", C2, "Análises por lote", "n.º", "const", D["analises_por_lote"], FMT_INT, "40 §5.2")
    p("d_cust", C2, "Custódia por lote", "€/lote", "const", D["inspecao_collateral_manager_eur_por_lote"], FMT_EUR, "[E, NV]")
    p("d_pag", C2, "Pagamentos bancários", "€/kg", "const", D["pagamentos_bancarios_eur_kg"], FMT_EKG, "[E]")
    C3 = "Azeite: canal granel certificado (só o Recomendado tem volume)"
    p("g_premio_venda", C3, "Preço de venda = índice + prémio de certificação", "€/kg", "const", G["premio_venda_sobre_indice_eur_kg"], FMT_EKG2,
      "[E] JSON: granel = índice + ≤ 0,15 €/kg, margem 0–0,10; 40 §1.3: margem bruta de trading 0,15–0,30. Usa-se 0,15")
    p("g_desc", C3, "Desconto ao comprador", "%", "const", G["desconto_efetivo"], FMT_PCT, "Venda ao índice + prémio")
    p("g_emb", C3, "Embalamento", "€/kg", "const", G["embalamento_eur_kg"], FMT_EKG2, "IBC/cisterna, sem embalamento")
    p("g_log", C3, "Logística (cisterna até ao embalador/exportador)", "€/kg", "const", G["logistica_eur_kg"], FMT_EKG2,
      "[E, NV] ~1.000 € por cisterna de 24 t na Ibéria")
    p("g_3pl", C3, "Armazém logístico (3PL)", "€/kg", "const", G["armazem_logistico_picking_eur_kg"], FMT_EKG2, "—")
    p("g_ds", C3, "Stock médio (back-to-back)", "dias", "const", G["dias_stock"], FMT_DAYS, "[E] lote sai do tanque em ~30 dias")
    p("g_dr", C3, "Prazo médio de recebimento", "dias", "const", G["dias_recebimento"], FMT_DAYS,
      "[E] igual aos distribuidores no base; Dir. 2019/633: ≤ 60 dias para não perecíveis")
    p("g_lote", C3, "Dimensão do lote", "t", "const", G["lote_medio_t"], FMT_T, "Igual à verificação para terceiros (lote de 100 t)")
    p("g_anal_n", C3, "Análises por lote", "n.º", "const", G["analises_por_lote"], FMT_INT, "600 € por lote de 100 t")
    p("g_cust", C3, "Custódia por lote", "€/lote", "const", G["inspecao_collateral_manager_eur_por_lote"], FMT_EUR, "1.200 € por lote de 100 t")
    p("g_pag", C3, "Pagamentos bancários", "€/kg", "const", G["pagamentos_bancarios_eur_kg"], FMT_EKG, "[E] poucas operações de valor alto")
    p("g_prov", C3, "Provisão de crédito", "% vendas", "const", G["provisao_credito_pct_vendas"], FMT_PCT2, "[E] igual aos distribuidores")
    C4 = "Azeite: custos físicos, custódia e cobertura"
    p("az_arm", C4, "Armazenagem no lagar", "€/t/mês", "const", cfc["armazenagem_eur_t_mes"], FMT_EKG2, "[E] método de custo, 2–8 €/t/mês. 40 §1.5")
    p("az_seg", C4, "Seguro do stock", "%/ano", "const", cfc["seguro_stock_pct_ano"], FMT_PCT2, "[E, NV] 0,1–0,3%/ano. 40 §4.5")
    p("az_anal_eur", C4, "Análise físico-química completa + painel", "€", "const", cfc["analise_fisico_quimica_painel_eur"], FMT_EUR, "[E, NV] 40 §5.2")
    p("az_perdas", C4, "Perdas físicas (uma vez)", "% valor", "const", cfc["perdas_fisicas_pct_valor"], FMT_PCT2, "[E] 40 §5.2")
    p("az_descl_prob", C4, "Probabilidade de desclassificação por trimestre", "%", "const", cfc["desclassificacao_prob_por_trimestre"], FMT_PCT,
      "[E] 3%/5%/15%/30% a 3/6/12/18 meses. 40 §5.2")
    p("az_descl_pen", C4, "Penalização por desclassificação", "€/kg", "const", cfc["desclassificacao_penalizacao_eur_kg"], FMT_EKG2,
      "[F] diferencial AOVE–lampante 0,24 (2026); 0,75 um ano antes")
    p("az_dp", C4, "Prazo de pagamento ao produtor", "dias", "const", cfc["dias_pagamento_produtor"], FMT_DAYS,
      "[F] Dir. (UE) 2019/633 e DL 118/2010 (DL 76/2021): ≤ 30/60 dias; política Sprout 30 dias")
    p("hedge_az", C4, "Custo explícito de cobertura do azeite", "% vendas", "const", P["custos"]["hedge"]["azeite"]["custo_explicito_pct"], FMT_PCT2,
      "[F] sem futuros de azeite: MFAO extinto (MIC XSRM expirado em jan. 2017); cobertura = back-to-back, indexação, offtake")
    C5 = "Azeite: mercado"
    p("az_cons_cliente", C5, "Consumo por cliente HoReCa/gourmet", "t/ano", "const", az["consumo_por_cliente_horeca_gourmet_t_ano"], FMT_EKG,
      "[E] 150–300 L/ano = 137–275 kg. 10 §5")
    p("az_prod_pt", C5, "Produção de azeite em Portugal (referência)", "kt", "const", az["producao_pt_referencia_kt"], FMT_T,
      "[E] ~175 kt (160 Olivum; 179 Governo). 00 §1.1")
    C6 = "Água (modelo agente; lote de comparação)"
    p("ag_l_pal", C6, "Litros por palete", "L", "const", ag["litros_por_palete"], FMT_INT, "[E] 500 garrafas de 1,5 L")
    p("ag_garrafas", C6, "Garrafas por palete", "n.º", "const", ag["garrafas_por_palete"], FMT_INT, "[E]")
    p("ag_ref_l", C6, "Preço B2B de referência", "€/L", "const", ag["preco_b2b_referencia_eur_l"], FMT_EKG2, "[E, NV] hipótese 0,30–0,60 €/L. 10 §2.3")
    p("ag_desc", C6, "Desconto ao comprador", "%", "const", ag["desconto_efetivo"], FMT_PCT, "Documento-fonte")
    p("ag_dr_com", C6, "Prazo de recebimento da comissão", "dias", "const", P["capital_de_giro"]["agua"]["dias_recebimento_comissao"], FMT_DAYS, "[E]")
    p("ag_exfab", C6, "Preço ex-fábrica médio (lote principal)", "€/L", "const", ag["preco_ex_fabrica_medio_eur_l"], FMT_EKG2,
      "[I] DGEG 2025: 479 M€ / ~2.000 M L")
    p("ag_sdr", C6, "Depósito SDR por embalagem < 3 L", "€", "const", ag["deposito_sdr_eur_por_embalagem"], FMT_EKG2, "[F] SDR desde 10-04-2026")
    p("ag_lote_pal", C6, "Lote-tipo de água", "paletes", "const", ag["lote_tipo"]["paletes"], FMT_INT, "01 §2.7")
    p("ag_desc_alt", C6, "Desconto alternativo no lote principal", "%", "const", ag["lote_tipo"]["desconto_alternativo_principal"], FMT_PCT, "01 §2.7")
    p("ag_troncal", C6, "Logística troncal por lote (principal)", "€/lote", "const", ag["lote_tipo"]["logistica_troncal_eur_por_lote"], FMT_EUR, "01 §2.7 [E]")
    p("ag_entrega", C6, "Entrega por litro (principal)", "€/L", "const", ag["lote_tipo"]["entrega_eur_l"], FMT_EKG2, "01 §2.7 [E]")
    p("ag_ops", C6, "Operação e pagamentos por lote (principal)", "€/lote", "const", ag["lote_tipo"]["operacao_pagamentos_eur_por_lote"], FMT_EUR, "01 §2.7 [E]")
    p("ag_dias_fin", C6, "Dias de financiamento do stock e do SDR (principal)", "dias", "const", ag["lote_tipo"]["dias_financiamento"], FMT_DAYS, "01 §2.7 [E]")
    C7 = "Café verde (principal; compra a importador UE)"
    p("cf_pc", C7, "Preço de compra (EUR, preço fixo, importador UE operador EUDR)", "€/kg", "const", cf["preco_compra_eur_kg"], FMT_EKG2,
      "[E, NV] KC 2,85 USD/lb = 5,46 €/kg a 1,15 + ~1,5 €/kg de diferencial, importação e logística")
    p("cf_pv", C7, "Preço de venda a microtorrefatores", "€/kg", "const", cf["preco_venda_eur_kg"], FMT_EKG2, "[E] NV, sem comparável verificado (margem bruta ~11%)")
    p("cf_log", C7, "Logística", "€/kg", "const", cf["logistica_eur_kg"], FMT_EKG2, "[E]")
    p("cf_ds", C7, "Stock médio", "dias", "const", cf["dias_stock"], FMT_DAYS, "[E]")
    p("cf_dr", C7, "Prazo de recebimento", "dias", "const", cf["dias_recebimento"], FMT_DAYS, "[E]")
    p("cf_dp", C7, "Prazo de pagamento ao importador", "dias", "const", cf["dias_pagamento"], FMT_DAYS, "[E]")
    p("cf_prov", C7, "Provisão de crédito", "% vendas", "const", cf["provisao_credito_pct_vendas"], FMT_PCT2, "[E]")
    p("cf_ops", C7, "Operações e pagamentos", "€/kg", "const", cf["pagamentos_ops_eur_kg"], FMT_EKG2, "[E]")
    C8 = "Serviços: verificação e financiamento de stock de terceiros"
    p("tp_taxa_util", C8, "Comissão de verificação por utilização", "% carteira", "const", ft["taxa_verificacao_por_utilizacao_pct"], FMT_PCT2,
      "[E] estruturação/verificação 1–2% do lote financiado. 50 §12.2")
    p("tp_rot", C8, "Rotações da carteira por ano", "n.º", "const", ft["rotacoes_por_ano"], FMT_NUM, "[E]")
    p("tp_taxa_mon", C8, "Comissão de monitorização anual", "% carteira", "const", ft["taxa_monitorizacao_anual_pct"], FMT_PCT2, "[E]")
    p("tp_custo_ver", C8, "Custo de verificação por kg verificado", "€/kg", "const", ft["custo_verificacao_eur_kg_verificado"], FMT_EKG,
      "[E] collateral manager 1.200 € + análises 600 € por lote de 100 t")
    p("tp_val_col", C8, "Valor do colateral", "€/kg", "const", ft["valor_colateral_eur_kg"], FMT_EKG2, "Índice do granel")
    p("tp_ltv", C8, "LTV do banco parceiro", "%", "const", ft["ltv"], FMT_PCT, "[E] 60–65%. 40 §5.3")
    C9 = "Financiamento: SPV (Sprout Capital, SA) e obrigações"
    p("fin_k_spv", C9, "Capital social da SPV (subordinado da 1.ª série)", "€", "const", fin["capital_social_spv_eur"], FMT_EUR,
      "[L] mínimo legal de SA 50.000 € (CSC art. 276.º/5); realizado em 2026")
    p("fin_res_pct_notes", C9, "Reserva de primeira perda: teto em % do saldo de obrigações", "%", "const",
      fin["reserva_primeira_perda"]["teto_pct_saldo_notes"], FMT_PCT, "40 §5.3 (P3)")
    p("fin_res_pct_margem", C9, "Reserva de primeira perda: % da margem de cada lote", "%", "const",
      fin["reserva_primeira_perda"]["pct_margem_lote"], FMT_PCT, "40 §5.3 (P3)")
    p("fin_ret_sing", C9, "Retenção sobre juros — particulares residentes (liberatória)", "%", "const",
      fin["retorno_investidor"]["retencao_singulares"], FMT_PCT, "[L] 28%; não reconfirmado online")
    p("fin_ticket", C9, "Ticket médio por investidor (estimativa)", "€", "const", fin["ticket_medio_investidor_eur"], FMT_EUR, "01 §7.3 [E]")
    p("fin_nocional", C9, "Custo nocional do capital próprio (memo; fora do EBITDA)", "%/ano", "const",
      ue["custo_nocional_capital_proprio_anual"], FMT_PCT, "01 §2.2")
    p("spv_setup", C9, "SPV: constituição (2027)", "€", "const", spv["setup_eur"], FMT_EUR, "[E] SA PT 15–40k. 20 §7.1")
    p("spv_anual", C9, "SPV: manutenção anual", "€/ano", "const", spv["anual_eur"], FMT_EUR, "[E] 10–30k/ano")
    p("spv_serie", C9, "SPV: custo por série emitida", "€/série", "const", spv["por_serie_eur"], FMT_EUR, "[E]")
    C10 = "Seguro de crédito, KYC e cobertura"
    p("sc_pct", C10, "Seguro de crédito: prémio sobre vendas HoReCa e distribuidores", "% vendas", "const", sc["premio_pct_volume_seguro"], FMT_PCT2,
      "[E, NV] 0,1–0,6%, mais alto na restauração. 40 §4.2")
    p("sc_pct_granel", C10, "Seguro de crédito: prémio sobre vendas de granel", "% vendas", "const", sc["premio_pct_volume_seguro_granel"], FMT_PCT2,
      "[E] compradores industriais; metade inferior do intervalo 0,1–0,6%")
    p("sc_desde", C10, "Seguro de crédito no opex desde", "ano", "const", sc["incluido_no_opex_desde"], FMT_YEAR, "01 §4.1")
    p("kyc_kyb", C10, "KYB por empresa", "€", "const", kyc["kyb_eur_por_empresa"], FMT_EKG2, "[L/NV] 3–10 €. 30 §4.2")
    p("kyc_screen", C10, "Screening por entidade e mês", "€/mês", "const", kyc["screening_eur_por_entidade_mes"], FMT_EKG2, "[L/NV] 30 §4.4")
    C11 = "Pessoal"
    p("tsu", C11, "TSU da entidade empregadora", "%", "const", pe["tsu_entidade_empregadora"], FMT_PCT2, "[L] Código Contributivo 23,75%")
    p("outros_enc", C11, "Outros encargos (acidentes de trabalho, medicina, FGCT)", "%", "const", pe["outros_encargos_pct"], FMT_PCT, "[E]")
    p("subs", C11, "Subsídio de refeição por FTE", "€/ano", "const", pe["subsidio_refeicao_eur_fte_ano"], FMT_EUR, "[L/NV] fora da TSU até aos limites legais")
    p("upd", C11, "Atualização salarial anual", "%", "const", pe["atualizacao_salarial_anual"], FMT_PCT, "[E]")
    p("ff_2027", C11, "Fundadores: fator salarial efetivo em 2027", "x", "const", pe["fator_salarial_fundadores_2027"]["fator_efetivo_2027"], FMT_NUM,
      "8/12 × 0,6 + 4/12 × 1 (60% até ao gate de 31-08-2027)")
    p("ff_gate", C11, "Fundadores: fator salarial até ao gate", "x", "const", pe["fator_salarial_fundadores_2027"]["fator_ate_gate"], FMT_NUM, "01 §5.5")
    p("gate_m", C11, "Meses até ao gate (jan.–ago. 2027)", "meses", "const", pe["fator_salarial_fundadores_2027"]["meses_ate_gate"], FMT_INT, "01 §5.1")
    for fn, v in pe["funcoes"].items():
        p("sal_" + fn, C11, f"Salário bruto anual (14 meses) — {ROLE_LABEL[fn]}", "€/ano", "const", v["salario_bruto_anual"], FMT_EUR,
          "[E, NV] mercado Lisboa/Porto 2026–27")
        p("var_" + fn, C11, f"Variável — {ROLE_LABEL[fn]}", "% salário", "const", v.get("variavel_pct", 0.0), FMT_PCT, "[E]")
    C12 = "Opex não salarial: parâmetros"
    p("cont", C12, "Contingência sobre as linhas de opex", "%", "const", ox["contingencia_pct"], FMT_PCT, "[E]")
    p("infl", C12, "Inflação anual do opex", "%", "const", ox["inflacao_anual"], FMT_PCT, "[E]")
    C13 = "Pré-arranque (4.º trimestre de 2026)"
    for k, v in P["custos"]["pre_arranque_2026_eur"].items():
        p("pre_" + k, C13, PRE_LABEL[k], "€", "const", v, FMT_EUR, "30 §10.1; 20 §7.1")
    C14 = "Capex e caixa"
    p("tel_eur", C14, "Telemetria por local (4–6 sensores radar + gateway)", "€/local", "const", cx["telemetria_eur_por_local"], FMT_EUR, "[E] 3.000–9.000 €. 30 §10.1")
    p("eq_eur", C14, "Equipamento informático por nova FTE", "€", "const", cx["equipamento_informatico_eur_por_nova_fte"], FMT_EUR, "[E]")
    p("vida_sw", C14, "Vida útil: software (plataforma)", "anos", "const", cx["vida_util_anos"]["software"], FMT_INT, "[L]")
    p("vida_eq", C14, "Vida útil: equipamento informático", "anos", "const", cx["vida_util_anos"]["equipamento_informatico"], FMT_INT, "[L]")
    p("vida_sens", C14, "Vida útil: sensores", "anos", "const", cx["vida_util_anos"]["sensores"], FMT_INT, "[L]")
    p("cmin", C14, "Caixa mínima operacional", "€", "const", P["capital_de_giro"]["caixa_minima_operacional_eur"], FMT_EUR,
      "[E] ~2,5 meses de opex de 2027")
    C15 = "Impostos [L] (confirmar com contabilista certificado)"
    p("irc_pme", C15, "IRC: taxa PME sobre os primeiros 50.000 €", "%", "const", imp["irc_taxa_pme_primeiros_50k"], FMT_PCT, "[L] não reconfirmado")
    p("irc_pme_lim", C15, "IRC: limite da matéria coletável à taxa PME", "€", "const", 50000, FMT_EUR, "[L]")
    p("derr_mun", C15, "Derrama municipal (máximo)", "% lucro tributável", "const", imp["derrama_municipal_max"], FMT_PCT2, "[L] RFAL, até 1,5%")
    ds_ = imp["derrama_estadual"]
    p("derr_e1_de", C15, "Derrama estadual: 1.º escalão a partir de", "€", "const", ds_[0]["de"], FMT_EUR, "[L] CIRC art. 87.º-A")
    p("derr_e1_tx", C15, "Derrama estadual: taxa do 1.º escalão", "%", "const", ds_[0]["taxa"], FMT_PCT, "[L]")
    p("derr_e2_de", C15, "Derrama estadual: 2.º escalão a partir de", "€", "const", ds_[1]["de"], FMT_EUR, "[L]")
    p("derr_e2_tx", C15, "Derrama estadual: taxa do 2.º escalão", "%", "const", ds_[1]["taxa"], FMT_PCT, "[L]")
    p("derr_e3_de", C15, "Derrama estadual: 3.º escalão a partir de", "€", "const", ds_[2]["de"], FMT_EUR, "[L]")
    p("derr_e3_tx", C15, "Derrama estadual: taxa do 3.º escalão", "%", "const", ds_[2]["taxa"], FMT_PCT, "[L]")
    p("reporte_lim", C15, "Reporte de prejuízos: dedução máxima", "% lucro tributável", "const",
      imp["reporte_prejuizos"]["limite_deducao_pct_lucro_tributavel"], FMT_PCT, "[L] OE2023, sem limite temporal")
    C16 = "Piloto, stress tests e unit economics"
    p("pil_h_t", C16, "Piloto: azeite HoReCa/gourmet", "t", "const", kp["volume"]["azeite_horeca_gourmet_t"], FMT_T, "01 §5.2")
    p("pil_d_t", C16, "Piloto: azeite para distribuidores com offtake", "t", "const", kp["volume"]["azeite_distribuidores_offtake_t"], FMT_T, "01 §5.2")
    p("s6_carry", C16, "Carry de 18 meses (% do valor do lote)", "%", "const", st["S6_detencao_18_meses"]["carry_pct"], FMT_PCT, "[E] 40 §5.2")
    p("s6_var_lo", C16, "VaR 95% de preço a 18 meses — volatilidade 23%", "%", "const", st["S6_detencao_18_meses"]["var95_preco_sigma23"], FMT_PCT, "[E] 40 §1.3")
    p("s6_var_hi", C16, "VaR 95% de preço a 18 meses — volatilidade 42%", "%", "const", st["S6_detencao_18_meses"]["var95_preco_sigma42"], FMT_PCT, "[E] 40 §1.3")
    p("exp_aberta", C16, "Exposição aberta máxima (piloto; política back-to-back)", "% book", "const", sens["exposicao_aberta_back_to_back"], FMT_PCT, "01 §4.1")
    p("ue_ano", C16, "Ano de regime dos lotes-tipo", "ano", "const", ue["ano_regime"], FMT_YEAR, "01 §2.1")
    C17 = "Sensibilidades e tese literal do documento-fonte"
    p("sens_desc_1", C17, "Sensibilidade: desconto HoReCa — caso 1", "%", "const", sens["desconto_horeca"][0], FMT_PCT, "Documento-fonte 10–15%")
    p("sens_desc_2", C17, "Sensibilidade: desconto HoReCa — caso 2", "%", "const", sens["desconto_horeca"][1], FMT_PCT, "")
    p("sens_desc_3", C17, "Sensibilidade: desconto HoReCa — caso 3", "%", "const", sens["desconto_horeca"][2], FMT_PCT, "")
    p("sens_vol_neg", C17, "Sensibilidade: volumes físicos — caso baixo", "%", "const", sens["volume"][0], FMT_PCT, "Estrutura fixa inalterada")
    p("sens_vol_pos", C17, "Sensibilidade: volumes físicos — caso alto", "%", "const", sens["volume"][1], FMT_PCT, "")
    p("sens_premio_topo", C17, "Caso combinado: prémio de origem pago por distribuidores/exportadores (topo)", "€/kg", "const",
      sens["premio_origem_distribuidores_topo"], FMT_EKG2, "Topo da condição 0,10–0,20 do 01 §5.4; com carteira de terceiros do otimista e volumes +30%")
    p("sens_preco", C17, "Sensibilidade: preço do AOVE (índice e referências de venda)", "%", "const", sens["preco_aove"], FMT_PCT,
      "Repasse proporcional; embalamento, logística e fixos inalterados")
    p("tese_m_min", C17, "Margem líquida do embalador — mínimo (2018/19)", "€/kg", "const", tese["margem_liquida_embalador_eur_kg"]["min"], FMT_EKG,
      "[F] Observatorio. 00 §1.4")
    p("tese_m_med", C17, "Margem líquida do embalador — média de 3 campanhas", "€/kg", "const", tese["margem_liquida_embalador_eur_kg"]["media_3_campanhas"], FMT_EKG,
      "[F→I] (0,187 + 0,499 + 0,206) / 3")
    p("tese_m_max", C17, "Margem líquida do embalador — máximo (2019/20)", "€/kg", "const", tese["margem_liquida_embalador_eur_kg"]["max"], FMT_EKG, "[F]")
    for i, v in enumerate(tese["captura_margem_embalador"]):
        p(f"tese_cap_{i + 1}", C17, f"Captura da margem do embalador — caso {i + 1}", "%", "const", v, FMT_PCT, "Documento-fonte 80–90%")

    # ------------------------------------------------------------ séries por cenário
    Q1 = "Azeite: volumes e clientes"
    p("h_cli", Q1, "Clientes HoReCa/gourmet ativos (média anual)", "n.º", "sseries", scen_vals(az["clientes_horeca_gourmet_ativos_medios"]), FMT_INT,
      "[E] SOM: 1.000–2.000 clientes no ano 3. 10 §5")
    p("h_vol", Q1, "Azeite HoReCa/gourmet (≈ clientes × 0,205 t)", "t", "sseries", scen_vals(az["volume_horeca_gourmet_t"]), FMT_T,
      "[E] limitado ao SOM; documento-fonte com valores arredondados")
    p("d_vol", Q1, "Azeite para distribuidores/exportação com offtake", "t", "sseries", scen_vals(az["volume_distribuidores_export_t"]), FMT_T,
      "[E] restante volume do plano (base 2031 = 1,3% da produção PT)")
    p("g_vol", Q1, "Azeite granel certificado", "t", "sseries", scen_vals(az["volume_granel_certificado_t"]), FMT_T,
      "[E] só no Recomendado: 3.500 t em 2031 = 2% da produção PT")
    Q2 = "Azeite: logística e stock"
    p("h_log", Q2, "Logística até à HoReCa", "€/kg", "sseries", scen_vals(H["logistica_eur_kg"]), FMT_EKG2, "[E] 0,15–0,30 €/kg. 10 §1.7")
    p("h_ds", Q2, "Stock médio HoReCa (180 dias com < 300 clientes; 45 depois)", "dias", "sseries", scen_vals(H["dias_stock"]), FMT_DAYS, "[E] 01 §2.8")
    p("d_log", Q2, "Logística FCA/DAP distribuidores", "€/kg", "sseries", scen_vals(D["logistica_eur_kg"]), FMT_EKG2, "[E]")
    Q3 = "Água, café, serviços"
    p("ag_pal", Q3, "Água (paletes de 750 L)", "paletes", "sseries", scen_vals(ag["volume_paletes"]), FMT_INT,
      "[E] pessimista: sai no gate; Recomendado: 0,9 paletes por cliente HoReCa (produto de cesto)")
    p("cf_vol", Q3, "Café verde (antes do ano de início conta 0)", "t", "sseries", scen_vals(cf["volume_t"]), FMT_T, "[E] 10–20 microtorrefatores")
    p("tp_cart", Q3, "Carteira média financiada por banco parceiro (verificada pela Sprout)", "€", "sseries", scen_vals(ft["carteira_media_eur"]), FMT_EUR,
      "[E] fase 2 (2028+); o banco é o mutuante (RGICSF)")
    p("dd", Q3, "Dados e índice (subscrições)", "€", "sseries", scen_vals(dd), FMT_EUR, "[E] conservador: Poolred e Mintec já publicam preços")
    Q4 = "Financiamento"
    p("fin_pct", Q4, "Capital de giro financiado por obrigações", "%", "sseries", scen_vals(fin["pct_capital_giro_azeite_financiado_por_notes"]), FMT_PCT,
      "[L] 0% em 2027: a SPV só emite 12 meses após o registo (CSC art. 348.º/2)")
    p("fin_cupao", Q4, "Cupão anual das obrigações (bruto)", "%", "sseries", scen_vals(fin["cupao_anual"]), FMT_PCT2,
      "[E] 8%/ano no modelo de carry; ECSP de retalho 10–13%/ano")
    p("spv_series", Q4, "Séries emitidas no ano", "n.º", "sseries", scen_vals(spv["series_por_ano"]), FMT_INT, "[E]")
    p("sc_min", Q4, "Seguro de crédito: prémio mínimo anual", "€", "sseries", scen_vals(sc["premio_minimo_anual_eur"]), FMT_EUR, "[E, NV]")
    Q5 = "Capex"
    p("tel_locais", Q5, "Novos locais com telemetria", "n.º", "sseries", scen_vals(cx["locais_novos"]), FMT_INT, "[E]")
    p("plat", Q5, "Desenvolvimento capitalizável da plataforma", "€", "sseries", scen_vals(cx["plataforma_desenvolvimento_capitalizavel_eur"]), FMT_EUR,
      "[E] reescrita clean-room de ledger, pagamentos e KYC. 30 §8.6")
    Q6 = "Pessoal (FTE por função)"
    for fn, v in pe["funcoes"].items():
        p("fte_" + fn, Q6, f"FTE — {ROLE_LABEL[fn]}", "FTE", "sseries", scen_vals(v["fte"]), FMT_FTE, ROLE_NOTE.get(fn, "[E]"))
    Q7 = "Opex não salarial (€ de 2027, antes de contingência e inflação)"
    for k in OPEX_KEYS:
        p("ox_" + k, Q7, OPEX_LABEL[k], "€", "sseries", scen_vals(ox[k]), FMT_EUR, OPEX_NOTE.get(k, "[E]"))

    # ------------------------------------------------------------ séries comuns
    R1 = "Séries comuns a todos os cenários"
    p("spv_indep", R1, "SPV: funções independentes (administrador, comité, representante comum)", "€", "cseries", spv["funcoes_independentes_eur"], FMT_EUR,
      "[E, NV] 01 §4.6")
    p("irc_rate", R1, "IRC: taxa geral", "%", "cseries", [imp["irc_taxa_geral"][str(y)] for y in YEARS], FMT_PCT, "[L] 18% 2027; 17% 2028+")
    return L


ROLE_LABEL = {
    "ceo_comercial": "CEO / comercial (fundador)", "cto": "CTO (fundador)", "engenheiro_software": "Engenheiro(a) de software",
    "risco_tesouraria": "Responsável de Risco e Tesouraria", "compliance_aml": "Compliance / AML", "operacoes_supply": "Operações e supply",
    "tecnico_qualidade_azeite": "Técnico(a) de qualidade do azeite", "gestor_conta_b2b": "Gestor(a) de conta B2B (~400 contas)",
    "controller_financeiro": "Controller financeiro", "analista_dados_produto": "Analista de dados/produto",
    "assistente_operacoes": "Assistente de operações", "trader_granel_exportacao": "Trader de granel e exportação",
}
ROLE_NOTE = {
    "risco_tesouraria": "Obrigatório antes da 1.ª emissão (01 §4.6); Recomendado: meio ano em 2027",
    "gestor_conta_b2b": "~400 contas ativas por gestor com self-service e SDD [I, NV]",
    "trader_granel_exportacao": "Só no Recomendado, desde 2029 [E, NV]",
    "engenheiro_software": "Recomendado: só depois do gate",
}
FOUNDERS = ["ceo_comercial", "cto"]
OPEX_KEYS = ["juridico_compliance_externo", "auditoria_roc", "contabilidade_externa", "tecnologia_cloud_saas_kyc", "seguranca_pentest_iso27001",
             "seguros_empresa_rc_do_cyber", "marketing_comercial_b2b", "escritorio", "viagens_deslocacoes", "formacao_recrutamento"]
OPEX_LABEL = {
    "juridico_compliance_externo": "Jurídico e compliance externo", "auditoria_roc": "Auditoria (ROC)", "contabilidade_externa": "Contabilidade externa",
    "tecnologia_cloud_saas_kyc": "Tecnologia: cloud, SaaS, KYC, faturação certificada", "seguranca_pentest_iso27001": "Segurança: pentest e ISO 27001",
    "seguros_empresa_rc_do_cyber": "Seguros da empresa (RC, D&O, ciber)", "marketing_comercial_b2b": "Marketing e comercial B2B",
    "escritorio": "Escritório", "viagens_deslocacoes": "Viagens e deslocações", "formacao_recrutamento": "Formação e recrutamento",
}
OPEX_NOTE = {
    "juridico_compliance_externo": "[E] 2027 inclui parecer independente, SPV e minutas (20 §11.14)",
    "seguranca_pentest_iso27001": "[E] pentest 5–15k; ISO 27001 15–50k no 1.º ano (30 §8.4)",
    "tecnologia_cloud_saas_kyc": "[E] Recomendado: metade das contas HoReCa do base",
    "marketing_comercial_b2b": "[E] Recomendado: aquisição HoReCa limitada; granel/exportação por relação direta",
    "escritorio": "[E] Recomendado: cowork até haver equipa",
    "contabilidade_externa": "[E] base: controller interno desde 2030; Recomendado: externa em todo o horizonte",
}
PRE_LABEL = {
    "juridico_parecer_constituicao_minutas": "Jurídico: parecer, constituição e minutas", "auditoria_backoffice_pentest": "Auditoria do backoffice e pentest",
    "aipd_rgpd": "AIPD (RGPD)", "integracoes_ledger_contratado": "Integrações e ledger contratado", "contingencia": "Contingência",
}


class PRegistry:
    """Posição de cada parâmetro na folha Pressupostos e acesso por cenário."""

    def __init__(self, params):
        self.by = {p.pid: p for p in params}
        self.params = params
        self.sel = "Pressupostos!$E$4"

    def ref(self, pid, scen, y=0):
        p = self.by[pid]
        if p.kind == "const":
            return f"Pressupostos!$C${p.row}"
        if p.kind == "scen":
            col = "G" if scen == 0 else "CDEF"[scen - 1]
            return f"Pressupostos!${col}${p.row}"
        if p.kind == "cseries":
            return f"Pressupostos!${YC[y]}${p.row}"
        r = p.rows[4] if scen == 0 else p.rows[scen - 1]
        return f"Pressupostos!${YC[y]}${r}"

    def row_range(self, pid, scen):
        """Intervalo C:G da série (para INDEX por ano)."""
        p = self.by[pid]
        if p.kind in ("const", "scen"):
            return None
        r = p.row if p.kind == "cseries" else (p.rows[4] if scen == 0 else p.rows[scen - 1])
        return f"Pressupostos!$C${r}:$G${r}"


def write_pressupostos(wb, params, P):
    ws = wb.create_sheet("Pressupostos")
    title(ws, "Pressupostos", "Todos os inputs do modelo (azul). Valores lidos de docs/business-plan/modelo/pressupostos.json em 13-10-2026. "
                              "Etiquetas: [F] facto com fonte; [L] referência legal a confirmar; [E] estimativa; [I] inferência; NV não verificado.")
    put(ws, "A4", "Cenário ativo (escolher na lista)", bold=True)
    put(ws, "C4", "Base", kind="input", bold=True, fill=FILL_SEL)
    dv = DataValidation(type="list", formula1='"Pessimista,Base,Otimista,Recomendado"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C4")
    put(ws, "D4", "índice", color=MUTED, italic=True)
    put(ws, "A5", "Muda as folhas Volumes, Receitas, Pessoal, Custos, DR, Cash flow e Capital & financiamento. Cenários, Unit economics e "
                  "Sensibilidade mostram cenários fixos (folha Motor).", color=MUTED, italic=True)
    widths = {"A": 62, "B": 14, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 95}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    r = 7
    # 1) escalares por cenário
    section_row(ws, r, "A. Pressupostos escalares por cenário")
    r += 1
    hdr_row = r
    header_row(ws, r, ["Parâmetro", "Unidade"] + SCEN_LABEL + ["Ativo", "Fonte / racional"])
    put(ws, "E4", f"=MATCH($C$4,$C${hdr_row}:$F${hdr_row},0)", fmt="0")
    r += 1
    sec = None
    for p in [x for x in params if x.kind == "scen"]:
        if p.sec != sec:
            sec = p.sec
            put(ws, f"A{r}", sec, bold=True, color=MUTED)
            r += 1
        p.row = r
        put(ws, f"A{r}", p.label)
        put(ws, f"B{r}", p.unit, color=MUTED)
        for i, s in enumerate(SCEN):
            put(ws, f"{'CDEF'[i]}{r}", p.vals[s], fmt=p.fmt, kind="input")
        put(ws, f"G{r}", f"=CHOOSE($E$4,C{r},D{r},E{r},F{r})", fmt=p.fmt)
        put(ws, f"H{r}", p.fonte, color=MUTED)
        r += 1
    r += 1
    # 2) constantes
    section_row(ws, r, "B. Constantes (iguais em todos os cenários)")
    r += 1
    header_row(ws, r, ["Parâmetro", "Unidade", "Valor", "", "", "", "Ativo", "Fonte / racional"])
    r += 1
    sec = None
    for p in [x for x in params if x.kind == "const"]:
        if p.sec != sec:
            sec = p.sec
            put(ws, f"A{r}", sec, bold=True, color=MUTED)
            r += 1
        p.row = r
        put(ws, f"A{r}", p.label)
        put(ws, f"B{r}", p.unit, color=MUTED)
        put(ws, f"C{r}", p.vals, fmt=p.fmt, kind="input")
        put(ws, f"G{r}", f"=C{r}", fmt=p.fmt)
        put(ws, f"H{r}", p.fonte, color=MUTED)
        r += 1
    r += 1
    # 3) séries por cenário
    section_row(ws, r, "C. Séries anuais por cenário")
    r += 1
    header_row(ws, r, ["Parâmetro / cenário", "Unidade"] + [str(y) for y in YEARS] + ["Fonte / racional"])
    r += 1
    sec = None
    for p in [x for x in params if x.kind == "sseries"]:
        if p.sec != sec:
            sec = p.sec
            put(ws, f"A{r}", sec, bold=True, color=MUTED)
            r += 1
        put(ws, f"A{r}", p.label, bold=True)
        put(ws, f"B{r}", p.unit, color=MUTED)
        put(ws, f"H{r}", p.fonte, color=MUTED)
        r += 1
        rows = []
        for i, s in enumerate(SCEN):
            put(ws, f"A{r}", "    " + SCEN_LABEL[i], color=MUTED)
            for y in range(NY):
                put(ws, f"{YC[y]}{r}", p.vals[s][y], fmt=p.fmt, kind="input")
            rows.append(r)
            r += 1
        put(ws, f"A{r}", "    Ativo", italic=True)
        for y in range(NY):
            c = YC[y]
            put(ws, f"{c}{r}", f"=CHOOSE($E$4,{c}{rows[0]},{c}{rows[1]},{c}{rows[2]},{c}{rows[3]})", fmt=p.fmt)
        rows.append(r)
        p.rows = rows
        r += 1
    r += 1
    section_row(ws, r, "D. Séries anuais comuns a todos os cenários")
    r += 1
    header_row(ws, r, ["Parâmetro", "Unidade"] + [str(y) for y in YEARS] + ["Fonte / racional"])
    r += 1
    for p in [x for x in params if x.kind == "cseries"]:
        p.row = r
        put(ws, f"A{r}", p.label)
        put(ws, f"B{r}", p.unit, color=MUTED)
        for y in range(NY):
            put(ws, f"{YC[y]}{r}", p.vals[y], fmt=p.fmt, kind="input")
        put(ws, f"H{r}", p.fonte, color=MUTED)
        r += 1
    ws.freeze_panes = "C9"
    return ws


# ----------------------------------------------------------------------------- motor: linhas
class Line:
    def __init__(self, lid, sheet, section, label, unit, fmt, fn, scalar=False, style=None, note=None):
        self.lid, self.sheet, self.section, self.label, self.unit = lid, sheet, section, label, unit
        self.fmt, self.fn, self.scalar, self.style, self.note = fmt, fn, scalar, style, note


CH = [("h", "HoReCa/gourmet"), ("d", "distribuidores/exportação"), ("g", "granel certificado")]


def define_lines():
    L = []

    def ln(lid, sheet, section, label, unit, fmt, fn, scalar=False, style=None, note=None):
        L.append(Line(lid, sheet, section, label, unit, fmt, fn, scalar, style, note))

    def ssum(c, y, terms):
        return "+".join(terms)

    # ---------------------------------------------------------------- Volumes
    V = "Volumes"
    ln("ano", V, "Calendário", "Ano", "", FMT_YEAR, lambda c, y: 2027 + y, style="year")
    s = "Azeite (t)"
    ln("v_h_cli", V, s, "Clientes HoReCa/gourmet ativos (média anual)", "n.º", FMT_INT, lambda c, y: c.p("h_cli", y))
    ln("v_h_t", V, s, "Azeite HoReCa/gourmet", "t", FMT_T, lambda c, y: c.p("h_vol", y))
    ln("v_d_t", V, s, "Azeite distribuidores/exportação (offtake)", "t", FMT_T, lambda c, y: c.p("d_vol", y))
    ln("v_g_t", V, s, "Azeite granel certificado", "t", FMT_T, lambda c, y: c.p("g_vol", y))
    ln("v_az_t", V, s, "Azeite total", "t", FMT_T, lambda c, y: f"{c.l('v_h_t', y)}+{c.l('v_d_t', y)}+{c.l('v_g_t', y)}", style="key")
    ln("v_h_mix", V, s, "Peso da HoReCa/gourmet no volume", "%", FMT_PCT,
       lambda c, y: f"IF({c.l('v_az_t', y)}>0,{c.l('v_h_t', y)}/{c.l('v_az_t', y)},0)")
    ln("v_quota", V, s, "Quota da produção portuguesa", "%", FMT_PCT2, lambda c, y: f"{c.l('v_az_t', y)}/({c.p('az_prod_pt')}*1000)")
    ln("v_lotes", V, s, "Lotes-tipo equivalentes", "n.º", FMT_T,
       lambda c, y: f"{c.l('v_h_t', y)}/{c.p('h_lote')}+{c.l('v_d_t', y)}/{c.p('d_lote')}+{c.l('v_g_t', y)}/{c.p('g_lote')}")
    s = "Água (modelo agente)"
    ln("v_ag_pal", V, s, "Água", "paletes", FMT_INT, lambda c, y: c.p("ag_pal", y))
    ln("v_ag_l", V, s, "Água", "L", FMT_INT, lambda c, y: f"{c.l('v_ag_pal', y)}*{c.p('ag_l_pal')}")
    s = "Café verde"
    ln("v_cf_t", V, s, "Café verde (0 antes do ano de início)", "t", FMT_T,
       lambda c, y: f"IF({c.l('ano', y)}<{c.p('cf_inicio')},0,{c.p('cf_vol', y)})")
    s = "Verificação para terceiros (banco parceiro)"
    ln("v_tp_cart", V, s, "Carteira média financiada pelo banco parceiro", "€", FMT_EUR, lambda c, y: c.p("tp_cart", y))
    ln("v_tp_t", V, s, "Azeite verificado para terceiros", "t", FMT_T,
       lambda c, y: f"{c.l('v_tp_cart', y)}/({c.p('tp_val_col')}*{c.p('tp_ltv')})*{c.p('tp_rot')}/1000")

    # ---------------------------------------------------------------- economia por kg (Custos)
    K = "Custos"
    for k, nm in CH:
        s = f"Economia por kg: {nm}"

        def p_(c, name, y, k=k):
            return c.p(f"{k}_{name}", y)

        ln(f"{k}_gr", K, s, "Preço ao produtor (índice + prémio de origem)", "€/kg", FMT_EKG,
           lambda c, y, k=k: f"{c.p('az_indice')}+{c.p(k + '_premio_lagar')}")
        if k == "h":
            fref = lambda c, y: c.p("h_ref")
        elif k == "d":
            fref = lambda c, y: f"{c.p('d_ref')}+{c.p('d_premio_venda')}"
        else:
            fref = lambda c, y: f"{c.p('az_indice')}+{c.p('g_premio_venda')}"
        ln(f"{k}_ref", K, s, "Preço de referência do comprador", "€/kg", FMT_EKG, fref)
        ln(f"{k}_price", K, s, "Preço de venda Sprout (referência − desconto)", "€/kg", FMT_EKG,
           lambda c, y, k=k: f"{c.l(k + '_ref', y)}*(1-{c.p(k + '_desc')})", style="key")
        ln(f"{k}_emb", K, s, "Embalamento", "€/kg", FMT_EKG, lambda c, y, k=k: c.p(k + "_emb"))
        ln(f"{k}_log", K, s, "Logística", "€/kg", FMT_EKG, lambda c, y, k=k: c.p(k + "_log", y))
        ln(f"{k}_3pl", K, s, "Armazém logístico (3PL) e picking", "€/kg", FMT_EKG, lambda c, y, k=k: c.p(k + "_3pl"))
        ln(f"{k}_ds", K, s, "Stock médio", "dias", FMT_DAYS, lambda c, y, k=k: c.p(k + "_ds", y))
        ln(f"{k}_dr", K, s, "Prazo médio de recebimento", "dias", FMT_DAYS, lambda c, y, k=k: c.p(k + "_dr"))
        ln(f"{k}_arm", K, s, "Armazenagem", "€/kg", FMT_EKG, lambda c, y, k=k: f"{c.p('az_arm')}*{c.l(k + '_ds', y)}/30/1000")
        ln(f"{k}_seg", K, s, "Seguro do stock", "€/kg", FMT_EKG, lambda c, y, k=k: f"{c.l(k + '_gr', y)}*{c.p('az_seg')}*{c.l(k + '_ds', y)}/365")
        ln(f"{k}_anal", K, s, "Análises (testes de qualidade)", "€/kg", FMT_EKG,
           lambda c, y, k=k: f"{c.p(k + '_anal_n')}*{c.p('az_anal_eur')}/({c.p(k + '_lote')}*1000)")
        ln(f"{k}_perdas", K, s, "Perdas físicas", "€/kg", FMT_EKG, lambda c, y, k=k: f"{c.p('az_perdas')}*{c.l(k + '_gr', y)}")
        ln(f"{k}_descl", K, s, "Desclassificação esperada", "€/kg", FMT_EKG,
           lambda c, y, k=k: f"{c.p('az_descl_prob')}*{c.p('az_descl_pen')}*{c.l(k + '_ds', y)}/90")
        ln(f"{k}_cust", K, s, "Custódia (collateral manager)", "€/kg", FMT_EKG, lambda c, y, k=k: f"{c.p(k + '_cust')}/({c.p(k + '_lote')}*1000)")
        ln(f"{k}_pag", K, s, "Pagamentos bancários", "€/kg", FMT_EKG, lambda c, y, k=k: c.p(k + "_pag"))
        ln(f"{k}_prov", K, s, "Provisão de crédito", "€/kg", FMT_EKG, lambda c, y, k=k: f"{c.p(k + '_prov')}*{c.l(k + '_price', y)}")
        ln(f"{k}_hedge", K, s, "Cobertura de preço explícita (sem futuros de azeite)", "€/kg", FMT_EKG,
           lambda c, y, k=k: f"{c.p('hedge_az')}*{c.l(k + '_price', y)}")
        ln(f"{k}_cupao", K, s, "Cupão das obrigações (financiamento do stock)", "€/kg", FMT_EKG,
           lambda c, y, k=k: (f"{c.p('fin_cupao', y)}*({c.l(k + '_gr', y)}+{c.l(k + '_emb', y)})*({c.l(k + '_ds', y)}+{c.l(k + '_dr', y)}-{c.p('az_dp')})/365"
                              f"*{c.p('fin_pct', y)}/{c.p('fin_util')}"))
        ln(f"{k}_custo", K, s, "Custo direto total", "€/kg", FMT_EKG,
           lambda c, y, k=k: "+".join(c.l(f"{k}_{x}", y) for x in ("gr", "emb", "log", "3pl", "arm", "seg", "anal", "perdas", "descl", "cust", "pag", "prov", "hedge", "cupao")))
        ln(f"{k}_margem", K, s, "Margem de contribuição", "€/kg", FMT_EKG, lambda c, y, k=k: f"{c.l(k + '_price', y)}-{c.l(k + '_custo', y)}", style="key")
        ln(f"{k}_margem_pct", K, s, "Margem de contribuição em % do preço", "%", FMT_PCT,
           lambda c, y, k=k: f"IF({c.l(k + '_price', y)}>0,{c.l(k + '_margem', y)}/{c.l(k + '_price', y)},0)")
        ln(f"{k}_cg", K, s, "Capital de giro por kg vendido", "€/kg", FMT_EKG,
           lambda c, y, k=k: (f"({c.l(k + '_gr', y)}+{c.l(k + '_emb', y)})*({c.l(k + '_ds', y)}-{c.p('az_dp')})/365"
                              f"+{c.l(k + '_price', y)}*{c.l(k + '_dr', y)}/365"))

    def vsum(c, y, expr):
        """Σ por canal: volume (t) × 1000 × expr(canal)."""
        return "+".join(f"{c.l('v_' + k + '_t', y)}*1000*({expr(k)})" for k, _ in CH)

    # ---------------------------------------------------------------- Receitas
    R = "Receitas"
    s = "Receitas por fluxo (€)"
    for k, nm in CH:
        ln(f"r_{k}", R, s, f"Vendas de azeite: {nm}", "€", FMT_EUR, lambda c, y, k=k: f"{c.l('v_' + k + '_t', y)}*1000*{c.l(k + '_price', y)}")
    ln("r_az", R, s, "Vendas de azeite", "€", FMT_EUR, lambda c, y: f"{c.l('r_h', y)}+{c.l('r_d', y)}+{c.l('r_g', y)}", style="sub")
    ln("r_cf", R, s, "Vendas de café verde", "€", FMT_EUR, lambda c, y: f"{c.l('v_cf_t', y)}*1000*{c.p('cf_pv')}")
    ln("gmv_ag", R, s, "Memo: GMV de água (faturado pelo engarrafador)", "€", FMT_EUR,
       lambda c, y: f"{c.l('v_ag_l', y)}*{c.p('ag_ref_l')}*(1-{c.p('ag_desc')})", style="memo")
    ln("r_ag", R, s, "Comissões de agente: água", "€", FMT_EUR, lambda c, y: f"{c.l('gmv_ag', y)}*{c.p('ag_com')}")
    ln("r_tp", R, s, "Comissões de originação e verificação (terceiros)", "€", FMT_EUR,
       lambda c, y: f"{c.l('v_tp_cart', y)}*({c.p('tp_taxa_util')}*{c.p('tp_rot')}+{c.p('tp_taxa_mon')})")
    ln("r_dd", R, s, "Dados e índice", "€", FMT_EUR, lambda c, y: c.p("dd", y))
    ln("r_total", R, s, "Receita total", "€", FMT_EUR,
       lambda c, y: "+".join(c.l(x, y) for x in ("r_az", "r_cf", "r_ag", "r_tp", "r_dd")), style="key")
    s = "GMV e capital de lotes (€)"
    ln("gmv_total", R, s, "GMV total (azeite + café + água)", "€", FMT_EUR, lambda c, y: f"{c.l('r_az', y)}+{c.l('r_cf', y)}+{c.l('gmv_ag', y)}", style="key")
    ln("lotes_valor", R, s, "Valor de compra dos lotes (granel + embalamento; café verde)", "€", FMT_EUR,
       lambda c, y: vsum(c, y, lambda k: f"{c.l(k + '_gr', y)}+{c.l(k + '_emb', y)}") + f"+{c.l('v_cf_t', y)}*1000*{c.p('cf_pc')}")
    ln("lotes_fin_obr", R, s, "dos quais financiados por obrigações (capital financiado)", "€", FMT_EUR,
       lambda c, y: f"{c.l('lotes_valor', y)}*{c.p('fin_pct', y)}")
    s = "Margem de contribuição por fluxo (€)"
    for k, nm in CH:
        ln(f"mc_{k}", R, s, f"Azeite: {nm}", "€", FMT_EUR, lambda c, y, k=k: f"{c.l('v_' + k + '_t', y)}*1000*{c.l(k + '_margem', y)}")
    ln("mc_az", R, s, "Azeite", "€", FMT_EUR, lambda c, y: f"{c.l('mc_h', y)}+{c.l('mc_d', y)}+{c.l('mc_g', y)}", style="sub")
    ln("mc_az_kg", R, s, "Margem média do azeite", "€/kg", FMT_EKG,
       lambda c, y: f"IF({c.l('v_az_t', y)}>0,{c.l('mc_az', y)}/({c.l('v_az_t', y)}*1000),0)")
    ln("mc_az_pct", R, s, "Margem do azeite em % das vendas de azeite", "%", FMT_PCT,
       lambda c, y: f"IF({c.l('r_az', y)}>0,{c.l('mc_az', y)}/{c.l('r_az', y)},0)")
    ln("cf_fin_kg", R, s, "Café: cupão por kg", "€/kg", FMT_EKG,
       lambda c, y: f"{c.p('fin_cupao', y)}*{c.p('cf_pc')}*({c.p('cf_ds')}+{c.p('cf_dr')}-{c.p('cf_dp')})/365*{c.p('fin_pct', y)}/{c.p('fin_util')}",
       style="memo")
    ln("mc_cf", R, s, "Café verde", "€", FMT_EUR,
       lambda c, y: (f"{c.l('v_cf_t', y)}*1000*({c.p('cf_pv')}-{c.p('cf_pc')}-{c.p('cf_log')}-{c.l('cf_fin_kg', y)}"
                     f"-{c.p('cf_prov')}*{c.p('cf_pv')}-{c.p('cf_ops')}-{c.p('cf_hedge')}*{c.p('cf_pc')})"))
    ln("mc_ag", R, s, "Água (comissão)", "€", FMT_EUR, lambda c, y: c.l("r_ag", y))
    ln("mc_tp", R, s, "Verificação para terceiros (líquida)", "€", FMT_EUR, lambda c, y: f"{c.l('r_tp', y)}-{c.l('cv_tp', y)}")
    ln("mc_dd", R, s, "Dados e índice", "€", FMT_EUR, lambda c, y: c.l("r_dd", y))
    ln("mc_total", R, s, "Margem de contribuição total", "€", FMT_EUR,
       lambda c, y: "+".join(c.l(x, y) for x in ("mc_az", "mc_cf", "mc_ag", "mc_tp", "mc_dd")), style="key")

    # ---------------------------------------------------------------- Custos totais
    s = "Custo das mercadorias e custos diretos (€)"
    ln("cmv_az", K, s, "Custo das mercadorias: azeite (granel + embalamento)", "€", FMT_EUR,
       lambda c, y: vsum(c, y, lambda k: f"{c.l(k + '_gr', y)}+{c.l(k + '_emb', y)}"))
    ln("cmv_cf", K, s, "Custo das mercadorias: café verde", "€", FMT_EUR, lambda c, y: f"{c.l('v_cf_t', y)}*1000*{c.p('cf_pc')}")
    ln("cmv", K, s, "Custo das mercadorias vendidas", "€", FMT_EUR, lambda c, y: f"{c.l('cmv_az', y)}+{c.l('cmv_cf', y)}", style="sub")
    ln("cv_log", K, s, "Logística", "€", FMT_EUR, lambda c, y: vsum(c, y, lambda k: c.l(k + "_log", y)) + f"+{c.l('v_cf_t', y)}*1000*{c.p('cf_log')}")
    ln("cv_3pl", K, s, "Armazém logístico (3PL) e picking", "€", FMT_EUR, lambda c, y: vsum(c, y, lambda k: c.l(k + "_3pl", y)))
    ln("cv_arm", K, s, "Armazenagem e seguro do stock", "€", FMT_EUR, lambda c, y: vsum(c, y, lambda k: f"{c.l(k + '_arm', y)}+{c.l(k + '_seg', y)}"))
    ln("cv_qual", K, s, "Análises, perdas e desclassificação", "€", FMT_EUR,
       lambda c, y: vsum(c, y, lambda k: f"{c.l(k + '_anal', y)}+{c.l(k + '_perdas', y)}+{c.l(k + '_descl', y)}"))
    ln("cv_cust", K, s, "Custódia (collateral manager)", "€", FMT_EUR, lambda c, y: vsum(c, y, lambda k: c.l(k + "_cust", y)))
    ln("cv_pag", K, s, "Pagamentos bancários e operações", "€", FMT_EUR,
       lambda c, y: vsum(c, y, lambda k: c.l(k + "_pag", y)) + f"+{c.l('v_cf_t', y)}*1000*{c.p('cf_ops')}")
    ln("cv_prov", K, s, "Provisão de crédito", "€", FMT_EUR,
       lambda c, y: vsum(c, y, lambda k: c.l(k + "_prov", y)) + f"+{c.l('v_cf_t', y)}*1000*{c.p('cf_prov')}*{c.p('cf_pv')}")
    ln("cv_hedge", K, s, "Cobertura explícita (café: forward EUR/USD; azeite: sem futuros)", "€", FMT_EUR,
       lambda c, y: vsum(c, y, lambda k: c.l(k + "_hedge", y)) + f"+{c.l('v_cf_t', y)}*1000*{c.p('cf_hedge')}*{c.p('cf_pc')}")
    ln("cv_tp", K, s, "Verificação para terceiros (collateral manager e análises)", "€", FMT_EUR,
       lambda c, y: f"{c.l('v_tp_t', y)}*1000*{c.p('tp_custo_ver')}*{c.p('tp_share')}")
    ln("cv_total", K, s, "Custos variáveis diretos", "€", FMT_EUR,
       lambda c, y: "+".join(c.l(x, y) for x in ("cv_log", "cv_3pl", "cv_arm", "cv_qual", "cv_cust", "cv_pag", "cv_prov", "cv_hedge", "cv_tp")), style="sub")
    ln("cup_az", K, s, "Cupão das obrigações: azeite", "€", FMT_EUR, lambda c, y: vsum(c, y, lambda k: c.l(k + "_cupao", y)))
    ln("cup_cf", K, s, "Cupão das obrigações: café", "€", FMT_EUR, lambda c, y: f"{c.l('v_cf_t', y)}*1000*{c.l('cf_fin_kg', y)}")
    ln("cup_total", K, s, "Cupão das obrigações (custo de financiar o stock)", "€", FMT_EUR, lambda c, y: f"{c.l('cup_az', y)}+{c.l('cup_cf', y)}", style="sub")
    ln("chk_mc", K, s, "Verificação: receita − CMV − variáveis − cupão − margem de contribuição", "€", FMT_EUR,
       lambda c, y: f"ROUND({c.l('r_total', y)}-{c.l('cmv', y)}-{c.l('cv_total', y)}-{c.l('cup_total', y)}-{c.l('mc_total', y)},6)", style="check")

    # ---------------------------------------------------------------- Pessoal
    PS = "Pessoal"
    roles = list(ROLE_LABEL.keys())
    s = "FTE por função"
    for fn in roles:
        ln("pf_" + fn, PS, s, ROLE_LABEL[fn], "FTE", FMT_FTE, lambda c, y, fn=fn: c.p("fte_" + fn, y))
    ln("fte_total", PS, s, "Total de FTE", "FTE", FMT_FTE, lambda c, y: "+".join(c.l("pf_" + fn, y) for fn in roles), style="key")
    s = "Custo salarial bruto (€, 14 meses, com variável e atualização)"
    for fn in roles:
        def fs(c, y, fn=fn):
            base = f"{c.l('pf_' + fn, y)}*{c.p('sal_' + fn)}*(1+{c.p('upd')})^{y}*(1+{c.p('var_' + fn)})"
            if fn in FOUNDERS and y == 0:
                base += f"*{c.p('ff_2027')}"
            return base
        ln("ps_" + fn, PS, s, ROLE_LABEL[fn], "€", FMT_EUR, fs)
    ln("ps_total", PS, s, "Salários brutos", "€", FMT_EUR, lambda c, y: "+".join(c.l("ps_" + fn, y) for fn in roles), style="sub")
    ln("ps_enc", PS, s, "TSU e outros encargos", "€", FMT_EUR, lambda c, y: f"{c.l('ps_total', y)}*({c.p('tsu')}+{c.p('outros_enc')})")
    ln("ps_sub", PS, s, "Subsídio de refeição", "€", FMT_EUR, lambda c, y: f"{c.l('fte_total', y)}*{c.p('subs')}*(1+{c.p('upd')})^{y}")
    ln("pess_total", PS, s, "Custo total com pessoal", "€", FMT_EUR, lambda c, y: f"{c.l('ps_total', y)}+{c.l('ps_enc', y)}+{c.l('ps_sub', y)}", style="key")
    ln("pess_fte", PS, s, "Custo médio por FTE", "€", FMT_EUR, lambda c, y: f"IF({c.l('fte_total', y)}>0,{c.l('pess_total', y)}/{c.l('fte_total', y)},0)", style="memo")

    # ---------------------------------------------------------------- Opex e capex (Custos)
    s = "Opex não salarial (€ nominais: + contingência e inflação)"
    for k in OPEX_KEYS:
        ln("ox_" + k, K, s, OPEX_LABEL[k], "€", FMT_EUR, lambda c, y, k=k: f"{c.p('ox_' + k, y)}*(1+{c.p('cont')})*(1+{c.p('infl')})^{y}")
    ln("ox_linhas", K, s, "Subtotal das linhas", "€", FMT_EUR, lambda c, y: "+".join(c.l("ox_" + k, y) for k in OPEX_KEYS), style="sub")
    ln("ox_spv", K, s, "SPV: constituição, manutenção, séries e funções independentes", "€", FMT_EUR,
       lambda c, y: (("" if y else f"{c.p('spv_setup')}+") + f"{c.p('spv_anual')}+{c.p('spv_serie')}*{c.p('spv_series', y)}+{c.p('spv_indep', y)}"))
    ln("ox_seg", K, s, "Seguro de crédito (prémio proporcional ou mínimo)", "€", FMT_EUR,
       lambda c, y: (f"IF({c.l('ano', y)}<{c.p('sc_desde')},0,MAX({c.p('sc_min', y)},"
                     f"{c.p('sc_pct')}*({c.l('r_h', y)}+{c.l('r_d', y)})+{c.p('sc_pct_granel')}*{c.l('r_g', y)}))"))
    ln("ox_total", K, s, "Opex não salarial", "€", FMT_EUR, lambda c, y: f"{c.l('ox_linhas', y)}+{c.l('ox_spv', y)}+{c.l('ox_seg', y)}", style="key")
    s = "Capex e amortizações (€)"
    ln("cx_tel", K, s, "Telemetria nos lagares", "€", FMT_EUR, lambda c, y: f"{c.p('tel_eur')}*{c.p('tel_locais', y)}")
    ln("cx_plat", K, s, "Plataforma (desenvolvimento capitalizável)", "€", FMT_EUR, lambda c, y: c.p("plat", y))
    ln("cx_eq", K, s, "Equipamento informático (novas FTE)", "€", FMT_EUR,
       lambda c, y: (f"{c.p('eq_eur')}*MAX(0,{c.l('fte_total', y)})" if y == 0 else f"{c.p('eq_eur')}*MAX(0,{c.l('fte_total', y)}-{c.l('fte_total', y - 1)})"))
    ln("capex", K, s, "Capex", "€", FMT_EUR, lambda c, y: f"{c.l('cx_tel', y)}+{c.l('cx_plat', y)}+{c.l('cx_eq', y)}", style="sub")

    def da(src, vida):
        def f(c, y):
            terms = "+".join(f"{c.l(src, j)}*(({y - j})<{c.p(vida)})" for j in range(y + 1))
            return f"({terms})/{c.p(vida)}"
        return f

    ln("da_plat", K, s, "Amortização: plataforma", "€", FMT_EUR, da("cx_plat", "vida_sw"))
    ln("da_tel", K, s, "Amortização: telemetria", "€", FMT_EUR, da("cx_tel", "vida_sens"))
    ln("da_eq", K, s, "Amortização: equipamento", "€", FMT_EUR, da("cx_eq", "vida_eq"))
    ln("da", K, s, "Amortizações", "€", FMT_EUR, lambda c, y: f"{c.l('da_plat', y)}+{c.l('da_tel', y)}+{c.l('da_eq', y)}", style="sub")

    # ---------------------------------------------------------------- DR
    DR = "DR"
    s = "Demonstração de resultados (€)"
    ln("dr_rec", DR, s, "Receita", "€", FMT_EUR, lambda c, y: c.l("r_total", y), style="key")
    ln("dr_cmv", DR, s, "Custo das mercadorias vendidas", "€", FMT_EUR, lambda c, y: f"-{c.l('cmv', y)}")
    ln("dr_mb", DR, s, "Margem bruta", "€", FMT_EUR, lambda c, y: f"{c.l('dr_rec', y)}+{c.l('dr_cmv', y)}", style="key")
    ln("dr_mb_pct", DR, s, "Margem bruta em % da receita", "%", FMT_PCT, lambda c, y: f"IF({c.l('dr_rec', y)}>0,{c.l('dr_mb', y)}/{c.l('dr_rec', y)},0)", style="memo")
    ln("dr_cv", DR, s, "Custos variáveis diretos (logística, custódia, qualidade, crédito, pagamentos)", "€", FMT_EUR, lambda c, y: f"-{c.l('cv_total', y)}")
    ln("dr_cup", DR, s, "Cupão das obrigações (custo de financiar o stock)", "€", FMT_EUR, lambda c, y: f"-{c.l('cup_total', y)}")
    ln("dr_mc", DR, s, "Margem de contribuição", "€", FMT_EUR, lambda c, y: f"{c.l('dr_mb', y)}+{c.l('dr_cv', y)}+{c.l('dr_cup', y)}", style="key")
    ln("dr_pess", DR, s, "Pessoal", "€", FMT_EUR, lambda c, y: f"-{c.l('pess_total', y)}")
    ln("dr_opex", DR, s, "Opex não salarial (incl. SPV e seguro de crédito)", "€", FMT_EUR, lambda c, y: f"-{c.l('ox_total', y)}")
    ln("dr_ebitda", DR, s, "EBITDA", "€", FMT_EUR, lambda c, y: f"{c.l('dr_mc', y)}+{c.l('dr_pess', y)}+{c.l('dr_opex', y)}", style="key",
       note="Convenção do modelo (igual ao 01 e a calculo_referencia.py): inclui o cupão, que financia stock lote a lote")
    ln("dr_ebitda_pct", DR, s, "EBITDA em % da receita", "%", FMT_PCT, lambda c, y: f"IF({c.l('dr_rec', y)}>0,{c.l('dr_ebitda', y)}/{c.l('dr_rec', y)},0)", style="memo")
    ln("dr_ebitda_xc", DR, s, "Memo: EBITDA sem o cupão (juros abaixo do EBIT, convenção SNC)", "€", FMT_EUR,
       lambda c, y: f"{c.l('dr_ebitda', y)}+{c.l('cup_total', y)}", style="memo")
    ln("dr_da", DR, s, "Amortizações", "€", FMT_EUR, lambda c, y: f"-{c.l('da', y)}")
    ln("dr_ebit", DR, s, "EBIT", "€", FMT_EUR, lambda c, y: f"{c.l('dr_ebitda', y)}+{c.l('dr_da', y)}", style="key")
    s = "IRC e derramas [L]"
    ln("tx_lt", DR, s, "Lucro tributável antes de prejuízos (= EBIT; sem correções fiscais)", "€", FMT_EUR, lambda c, y: c.l("dr_ebit", y))
    ln("tx_prej_ini", DR, s, "Prejuízos fiscais reportáveis no início do ano", "€", FMT_EUR, lambda c, y: "0" if y == 0 else c.l("tx_prej_fim", y - 1))
    ln("tx_prej_ger", DR, s, "Prejuízo fiscal gerado no ano", "€", FMT_EUR, lambda c, y: f"MAX(0,-{c.l('tx_lt', y)})")
    ln("tx_ded", DR, s, "Dedução de prejuízos (até 65% do lucro tributável)", "€", FMT_EUR,
       lambda c, y: f"MIN({c.l('tx_prej_ini', y)},{c.p('reporte_lim')}*MAX(0,{c.l('tx_lt', y)}))")
    ln("tx_prej_fim", DR, s, "Prejuízos fiscais reportáveis no fim do ano", "€", FMT_EUR,
       lambda c, y: f"{c.l('tx_prej_ini', y)}+{c.l('tx_prej_ger', y)}-{c.l('tx_ded', y)}")
    ln("tx_mc", DR, s, "Matéria coletável", "€", FMT_EUR, lambda c, y: f"MAX(0,{c.l('tx_lt', y)})-{c.l('tx_ded', y)}")
    ln("tx_irc", DR, s, "IRC (15% até 50.000 €; taxa geral acima)", "€", FMT_EUR,
       lambda c, y: f"{c.p('irc_pme')}*MIN({c.l('tx_mc', y)},{c.p('irc_pme_lim')})+{c.p('irc_rate', y)}*MAX(0,{c.l('tx_mc', y)}-{c.p('irc_pme_lim')})")
    ln("tx_dm", DR, s, "Derrama municipal", "€", FMT_EUR, lambda c, y: f"{c.p('derr_mun')}*MAX(0,{c.l('tx_lt', y)})")
    ln("tx_de", DR, s, "Derrama estadual", "€", FMT_EUR,
       lambda c, y: (f"{c.p('derr_e1_tx')}*MAX(0,MIN({c.l('tx_lt', y)},{c.p('derr_e2_de')})-{c.p('derr_e1_de')})"
                     f"+{c.p('derr_e2_tx')}*MAX(0,MIN({c.l('tx_lt', y)},{c.p('derr_e3_de')})-{c.p('derr_e2_de')})"
                     f"+{c.p('derr_e3_tx')}*MAX(0,{c.l('tx_lt', y)}-{c.p('derr_e3_de')})"))
    ln("tx_total", DR, s, "Imposto sobre o rendimento", "€", FMT_EUR, lambda c, y: f"{c.l('tx_irc', y)}+{c.l('tx_dm', y)}+{c.l('tx_de', y)}", style="sub")
    s = "Resultado"
    ln("dr_imp", DR, s, "Imposto sobre o rendimento", "€", FMT_EUR, lambda c, y: f"-{c.l('tx_total', y)}")
    ln("dr_rl", DR, s, "Resultado líquido", "€", FMT_EUR, lambda c, y: f"{c.l('dr_ebit', y)}+{c.l('dr_imp', y)}", style="key")

    # ---------------------------------------------------------------- Cash flow
    CF = "Cash flow"
    s = "Capital de giro (€, nível médio do ano, sem IVA)"
    ln("cg_stock_az", CF, s, "Stock de azeite (ao custo)", "€", FMT_EUR,
       lambda c, y: vsum(c, y, lambda k: f"({c.l(k + '_gr', y)}+{c.l(k + '_emb', y)})*{c.l(k + '_ds', y)}/365"))
    ln("cg_stock", CF, s, "Stock total (azeite + café)", "€", FMT_EUR,
       lambda c, y: f"{c.l('cg_stock_az', y)}+{c.l('v_cf_t', y)}*1000*{c.p('cf_pc')}*{c.p('cf_ds')}/365")
    ln("cg_rec", CF, s, "Recebíveis (prazos médios de recebimento)", "€", FMT_EUR,
       lambda c, y: (vsum(c, y, lambda k: f"{c.l(k + '_price', y)}*{c.l(k + '_dr', y)}/365")
                     + f"+{c.l('r_cf', y)}*{c.p('cf_dr')}/365+{c.l('r_ag', y)}*{c.p('ag_dr_com')}/365"))
    ln("cg_forn", CF, s, "Fornecedores (produtor pago a 30 dias)", "€", FMT_EUR,
       lambda c, y: (vsum(c, y, lambda k: f"({c.l(k + '_gr', y)}+{c.l(k + '_emb', y)})*{c.p('az_dp')}/365")
                     + f"+{c.l('v_cf_t', y)}*1000*{c.p('cf_pc')}*{c.p('cf_dp')}/365"))
    ln("cg", CF, s, "Capital de giro líquido", "€", FMT_EUR, lambda c, y: f"{c.l('cg_stock', y)}+{c.l('cg_rec', y)}-{c.l('cg_forn', y)}", style="key")
    s = "Financiamento do capital de giro (€)"
    ln("ob", CF, s, "Obrigações em circulação (capital levantado nas séries)", "€", FMT_EUR,
       lambda c, y: f"{c.p('fin_pct', y)}*{c.l('cg', y)}/{c.p('fin_util')}", style="key")
    ln("ob_uso", CF, s, "das quais aplicadas em stock e recebíveis", "€", FMT_EUR, lambda c, y: f"{c.p('fin_pct', y)}*{c.l('cg', y)}")
    ln("ob_float", CF, s, "Caixa parada nas séries (float de financiamento)", "€", FMT_EUR, lambda c, y: f"{c.l('ob', y)}-{c.l('ob_uso', y)}")
    ln("cg_eq", CF, s, "Capital de giro financiado por capital próprio", "€", FMT_EUR, lambda c, y: f"{c.l('cg', y)}-{c.l('ob_uso', y)}")
    ln("spv_caixa", CF, s, "Capital social da SPV ainda em caixa (não aplicado em stock)", "€", FMT_EUR,
       lambda c, y: f"MAX(0,{c.p('fin_k_spv')}-{c.l('cg_eq', y)})")
    ln("res", CF, s, "Reserva de primeira perda (caixa restrita, 5% das obrigações)", "€", FMT_EUR, lambda c, y: f"{c.p('fin_res_pct_notes')}*{c.l('ob', y)}")
    ln("inv_n", CF, s, "Investidores estimados (obrigações ÷ ticket médio)", "n.º", FMT_INT, lambda c, y: f"{c.l('ob', y)}/{c.p('fin_ticket')}", style="memo")
    s = "Cash flow (€)"
    ln("cf_ebitda", CF, s, "EBITDA", "€", FMT_EUR, lambda c, y: c.l("dr_ebitda", y))
    ln("cf_imp", CF, s, "Imposto pago (no próprio ano; simplificação)", "€", FMT_EUR, lambda c, y: f"-{c.l('tx_total', y)}")
    ln("cf_dcg", CF, s, "Variação do capital de giro", "€", FMT_EUR, lambda c, y: f"-{c.l('cg', y)}" if y == 0 else f"-({c.l('cg', y)}-{c.l('cg', y - 1)})")
    ln("cf_capex", CF, s, "Capex", "€", FMT_EUR, lambda c, y: f"-{c.l('capex', y)}")
    ln("cf_fcf", CF, s, "Cash flow operacional e de investimento", "€", FMT_EUR,
       lambda c, y: "+".join(c.l(x, y) for x in ("cf_ebitda", "cf_imp", "cf_dcg", "cf_capex")), style="sub")
    ln("cf_dob", CF, s, "Emissão líquida de obrigações", "€", FMT_EUR, lambda c, y: c.l("ob", y) if y == 0 else f"{c.l('ob', y)}-{c.l('ob', y - 1)}")
    ln("cf_dfloat", CF, s, "Variação da caixa parada nas séries", "€", FMT_EUR,
       lambda c, y: f"-{c.l('ob_float', y)}" if y == 0 else f"-({c.l('ob_float', y)}-{c.l('ob_float', y - 1)})")
    ln("cf_dres", CF, s, "Variação da reserva de primeira perda", "€", FMT_EUR,
       lambda c, y: f"-{c.l('res', y)}" if y == 0 else f"-({c.l('res', y)}-{c.l('res', y - 1)})")
    ln("cf_dspv", CF, s, "Capital social da SPV aplicado em stock", "€", FMT_EUR,
       lambda c, y: f"-({c.l('spv_caixa', y)}-{c.p('fin_k_spv')})" if y == 0 else f"-({c.l('spv_caixa', y)}-{c.l('spv_caixa', y - 1)})")
    ln("cf_acion", CF, s, "Cash flow para o acionista", "€", FMT_EUR,
       lambda c, y: "+".join(c.l(x, y) for x in ("cf_fcf", "cf_dob", "cf_dfloat", "cf_dres", "cf_dspv")), style="key")
    ln("cf_acum", CF, s, "Cash flow para o acionista acumulado", "€", FMT_EUR,
       lambda c, y: c.l("cf_acion", 0) if y == 0 else f"{c.l('cf_acum', y - 1)}+{c.l('cf_acion', y)}", style="key")
    ln("chk_cf", CF, s, "Verificação: EBITDA − capex − imposto − Δ max(capital próprio no giro; capital SPV) − Δ reserva − cash flow", "€", FMT_EUR,
       lambda c, y: (f"ROUND({c.l('dr_ebitda', y)}-{c.l('capex', y)}-{c.l('tx_total', y)}"
                     f"-(MAX({c.l('cg_eq', y)},{c.p('fin_k_spv')})-" + (f"{c.p('fin_k_spv')}" if y == 0 else f"MAX({c.l('cg_eq', y - 1)},{c.p('fin_k_spv')})") + ")"
                     f"-({c.l('res', y)}" + ("" if y == 0 else f"-{c.l('res', y - 1)}") + f")-{c.l('cf_acion', y)},6)"), style="check")
    s = "Caixa e rondas de capital (€)"
    ln("cx_ini", CF, s, "Caixa livre no início do ano", "€", FMT_EUR, lambda c, y: "0" if y == 0 else c.l("cx_fim", y - 1))
    ln("cx_pre", CF, s, "Pré-arranque (4.º trim. 2026) e capital social da SPV", "€", FMT_EUR,
       lambda c, y: ("-(" + "+".join(c.p(x) for x in PRE_IDS) + f"+{c.p('fin_k_spv')})") if y == 0 else "0")
    ln("cx_seed", CF, s, "Ronda pré-seed", "€", FMT_EUR, lambda c, y: c.p("fin_pre_seed") if y == 0 else "0")
    ln("cx_antes", CF, s, "Caixa antes de rondas adicionais", "€", FMT_EUR,
       lambda c, y: "+".join(c.l(x, y) for x in ("cx_ini", "cx_pre", "cx_seed", "cf_acion")))
    ln("cx_plug", CF, s, "Ronda adicional necessária para manter a caixa mínima", "€", FMT_EUR,
       lambda c, y: f"MAX(0,{c.p('cmin')}-{c.l('cx_antes', y)})", style="key")
    ln("cx_fim", CF, s, "Caixa livre no fim do ano", "€", FMT_EUR, lambda c, y: f"{c.l('cx_antes', y)}+{c.l('cx_plug', y)}", style="key")
    ln("cx_restrita", CF, s, "Caixa restrita (séries, reserva, capital da SPV)", "€", FMT_EUR,
       lambda c, y: f"{c.l('ob_float', y)}+{c.l('res', y)}+{c.l('spv_caixa', y)}")
    ln("cx_total", CF, s, "Caixa total", "€", FMT_EUR, lambda c, y: f"{c.l('cx_fim', y)}+{c.l('cx_restrita', y)}")
    ln("cx_rondas_acum", CF, s, "Capital próprio levantado (acumulado)", "€", FMT_EUR,
       lambda c, y: (f"{c.l('cx_seed', y)}+{c.l('cx_plug', y)}" if y == 0 else f"{c.l('cx_rondas_acum', y - 1)}+{c.l('cx_seed', y)}+{c.l('cx_plug', y)}"))

    # ---------------------------------------------------------------- Capital & financiamento
    KF = "Capital & financiamento"

    def rng(c, lid):
        return c.rng(lid)

    s = "Necessidade de capital próprio"
    ln("k_pre", KF, s, "Pré-arranque (4.º trim. 2026)", "€", FMT_EUR, lambda c, y: "+".join(c.p(x) for x in PRE_IDS), scalar=True)
    ln("k_spv", KF, s, "Capital social da SPV", "€", FMT_EUR, lambda c, y: c.p("fin_k_spv"), scalar=True)
    ln("k_cmin", KF, s, "Caixa mínima operacional", "€", FMT_EUR, lambda c, y: c.p("cmin"), scalar=True)
    ln("k_vale", KF, s, "Vale do cash flow acumulado para o acionista", "€", FMT_EUR, lambda c, y: f"MAX(0,-MIN({rng(c, 'cf_acum')}))", scalar=True)
    ln("k_need", KF, s, "Necessidade de capital próprio (pico no horizonte)", "€", FMT_EUR,
       lambda c, y: f"{c.l('k_pre', 0)}+{c.l('k_spv', 0)}+{c.l('k_cmin', 0)}+{c.l('k_vale', 0)}", scalar=True, style="key")
    ln("k_min_ano", KF, s, "Ano do mínimo do cash flow acumulado", "ano", FMT_YEAR,
       lambda c, y: f"INDEX({rng(c, 'ano')},1,MATCH(MIN({rng(c, 'cf_acum')}),{rng(c, 'cf_acum')},0))", scalar=True)
    ln("k_ainda", KF, s, "O acumulado ainda desce em 2031?", "", None,
       lambda c, y: f"IF({c.l('cf_acion', 4)}<0,\"sim\",\"não\")", scalar=True)
    ln("k_rondas", KF, s, "Capital próprio levantado no plano (pré-seed + rondas adicionais)", "€", FMT_EUR,
       lambda c, y: c.l("cx_rondas_acum", 4), scalar=True)
    s = "Aplicação de fundos até ao ano do mínimo (€)"
    ln("u_pre", KF, s, "Pré-arranque e capital social da SPV", "€", FMT_EUR, lambda c, y: f"{c.l('k_pre', 0)}+{c.l('k_spv', 0)}", scalar=True)
    ln("u_ebitda", KF, s, "Perdas operacionais (EBITDA negativo acumulado)", "€", FMT_EUR,
       lambda c, y: f"-SUMPRODUCT(({rng(c, 'ano')}<={c.l('k_min_ano', 0)})*{rng(c, 'dr_ebitda')})", scalar=True)
    ln("u_capex", KF, s, "Capex", "€", FMT_EUR, lambda c, y: f"SUMPRODUCT(({rng(c, 'ano')}<={c.l('k_min_ano', 0)})*{rng(c, 'capex')})", scalar=True)
    ln("u_imp", KF, s, "Impostos", "€", FMT_EUR, lambda c, y: f"SUMPRODUCT(({rng(c, 'ano')}<={c.l('k_min_ano', 0)})*{rng(c, 'tx_total')})", scalar=True)
    ln("u_cg", KF, s, "Capital de giro, reserva e caixa das séries a cargo do capital próprio", "€", FMT_EUR,
       lambda c, y: f"{c.l('k_vale', 0)}-{c.l('u_ebitda', 0)}-{c.l('u_capex', 0)}-{c.l('u_imp', 0)}", scalar=True)
    ln("u_cmin", KF, s, "Caixa mínima operacional", "€", FMT_EUR, lambda c, y: c.l("k_cmin", 0), scalar=True)
    ln("u_total", KF, s, "Total = necessidade de capital próprio", "€", FMT_EUR,
       lambda c, y: "+".join(c.l(x, 0) for x in ("u_pre", "u_ebitda", "u_capex", "u_imp", "u_cg", "u_cmin")), scalar=True, style="key")
    s = "Pré-seed, runway e gate"
    ln("b_seed", KF, s, "Ronda pré-seed", "€", FMT_EUR, lambda c, y: c.p("fin_pre_seed"), scalar=True)
    ln("b_gate", KF, s, "Gasto até ao gate de 31-08-2027 (pré-arranque, SPV, opex jan.–ago., capex, stock do piloto)", "€", FMT_EUR,
       lambda c, y: (f"{c.l('k_pre', 0)}+{c.p('fin_k_spv')}+{c.p('gate_m')}/12*(("
                     f"{c.l('ps_total', 0)}-{c.l('ps_ceo_comercial', 0)}-{c.l('ps_cto', 0)}+({c.l('ps_ceo_comercial', 0)}+{c.l('ps_cto', 0)})/{c.p('ff_2027')}*{c.p('ff_gate')})"
                     f"*(1+{c.p('tsu')}+{c.p('outros_enc')})+{c.l('fte_total', 0)}*{c.p('subs')}+{c.l('ox_total', 0)})"
                     f"+{c.l('capex', 0)}+MAX(0,{c.p('pil_h_t')}*1000*({c.l('h_gr', 0)}+{c.l('h_emb', 0)})+{c.p('pil_d_t')}*1000*({c.l('d_gr', 0)}+{c.l('d_emb', 0)})-{c.p('fin_k_spv')})"),
       scalar=True, note="01 §5.5: ~0,74 M€ no base")
    ln("b_2027", KF, s, "Gasto até 31-12-2027 (pré-arranque, SPV e cash flow de 2027)", "€", FMT_EUR,
       lambda c, y: f"{c.l('k_pre', 0)}+{c.p('fin_k_spv')}-{c.l('cf_acion', 0)}", scalar=True)
    ln("rw_cash", KF, s, "Caixa sem rondas adicionais (pré-seed − pré-arranque − SPV + cash flow acumulado)", "€", FMT_EUR,
       lambda c, y: f"{c.p('fin_pre_seed')}-{c.l('k_pre', 0)}-{c.p('fin_k_spv')}+{c.l('cf_acum', y)}")
    ln("rw_cross", KF, s, "Mês em que a caixa sem rondas adicionais chega a zero (1 = jan. 2027)", "mês", FMT_MES,
       lambda c, y: ((f"IF(AND({c.p('fin_pre_seed')}-{c.l('k_pre', 0)}-{c.p('fin_k_spv')}>0,{c.l('rw_cash', 0)}<=0),"
                      f"12*({c.p('fin_pre_seed')}-{c.l('k_pre', 0)}-{c.p('fin_k_spv')})/({c.p('fin_pre_seed')}-{c.l('k_pre', 0)}-{c.p('fin_k_spv')}-{c.l('rw_cash', 0)}),\"\")")
                     if y == 0 else
                     f"IF(AND({c.l('rw_cash', y - 1)}>0,{c.l('rw_cash', y)}<=0),{12 * y}+12*{c.l('rw_cash', y - 1)}/({c.l('rw_cash', y - 1)}-{c.l('rw_cash', y)}),\"\")"))
    ln("b_runway", KF, s, "Runway da pré-seed (meses desde jan. 2027; sem rondas adicionais)", "meses", FMT_MES,
       lambda c, y: f"IF(COUNT({rng(c, 'rw_cross')})=0,\"> 60\",MIN({rng(c, 'rw_cross')}))", scalar=True, style="key")
    s = "Break-even"
    ln("be_cross", KF, s, "Cruzamento do EBITDA mensal (run-rate) com zero (mês; 1 = jan. 2027)", "mês", FMT_MES,
       lambda c, y: (f"IF({c.l('dr_ebitda', 0)}>=0,1,\"\")" if y == 0 else
                     f"IF(AND({c.l('dr_ebitda', y - 1)}<0,{c.l('dr_ebitda', y)}>=0),{6.5 + 12 * (y - 1)}+12*(-{c.l('dr_ebitda', y - 1)})/({c.l('dr_ebitda', y)}-{c.l('dr_ebitda', y - 1)}),\"\")"))
    ln("be_mes", KF, s, "Mês de break-even do EBITDA (run-rate; 1 = jan. 2027)", "mês", FMT_MES,
       lambda c, y: f"IF(COUNT({rng(c, 'be_cross')})=0,\"não atinge até 2031\",ROUNDUP(MIN({rng(c, 'be_cross')}),0))", scalar=True)
    ln("be_data", KF, s, "Break-even do EBITDA (mês)", "data", FMT_DATE,
       lambda c, y: f"IF(ISNUMBER({c.l('be_mes', 0)}),DATE(2027,{c.l('be_mes', 0)},1),\"não atinge até 2031\")", scalar=True, style="key")
    ln("be_f_ebitda", KF, s, "Anos com EBITDA ≥ 0", "ano", FMT_YEAR, lambda c, y: f"IF({c.l('dr_ebitda', y)}>=0,{c.l('ano', y)},\"\")")
    ln("be_ano_ebitda", KF, s, "1.º ano com EBITDA ≥ 0", "ano", FMT_YEAR,
       lambda c, y: f"IF(COUNT({rng(c, 'be_f_ebitda')})=0,\"não atinge até 2031\",MIN({rng(c, 'be_f_ebitda')}))", scalar=True)
    ln("be_f_rl", KF, s, "Anos com resultado líquido ≥ 0", "ano", FMT_YEAR, lambda c, y: f"IF({c.l('dr_rl', y)}>=0,{c.l('ano', y)},\"\")")
    ln("be_ano_rl", KF, s, "1.º ano com resultado líquido ≥ 0", "ano", FMT_YEAR,
       lambda c, y: f"IF(COUNT({rng(c, 'be_f_rl')})=0,\"não atinge até 2031\",MIN({rng(c, 'be_f_rl')}))", scalar=True)
    ln("be_f_cf", KF, s, "Anos com cash flow para o acionista ≥ 0", "ano", FMT_YEAR, lambda c, y: f"IF({c.l('cf_acion', y)}>=0,{c.l('ano', y)},\"\")")
    ln("be_ano_cf", KF, s, "1.º ano com cash flow para o acionista ≥ 0", "ano", FMT_YEAR,
       lambda c, y: f"IF(COUNT({rng(c, 'be_f_cf')})=0,\"não atinge até 2031\",MIN({rng(c, 'be_f_cf')}))", scalar=True)
    s = "Investidores (obrigações da Sprout Capital, SA)"
    ln("i_ob", KF, s, "Obrigações em circulação", "€", FMT_EUR, lambda c, y: c.l("ob", y))
    ln("i_cup", KF, s, "Cupão bruto anual", "%", FMT_PCT2, lambda c, y: c.p("fin_cupao", y))
    ln("i_cup_liq", KF, s, "Cupão líquido para um particular residente (retenção de 28%)", "%", FMT_PCT2,
       lambda c, y: f"{c.p('fin_cupao', y)}*(1-{c.p('fin_ret_sing')})")
    ln("i_pago", KF, s, "Cupão pago no ano", "€", FMT_EUR, lambda c, y: c.l("cup_total", y))
    ln("i_alm", KF, s, "Almofada de primeira perda (capital próprio no giro + reserva) em % das obrigações", "%", FMT_PCT,
       lambda c, y: f"IF({c.l('ob', y)}>0,({c.l('cg_eq', y)}+{c.l('res', y)})/{c.l('ob', y)},0)")
    ln("i_n", KF, s, "Investidores estimados", "n.º", FMT_INT, lambda c, y: c.l("inv_n", y))
    s = "O que tem de ser verdade para EBITDA = 0 (uma alavanca de cada vez)"
    ln("wt_gap", KF, s, "Défice de EBITDA a cobrir", "€", FMT_EUR, lambda c, y: f"MAX(0,-{c.l('dr_ebitda', y)})")
    ln("wt_ins", KF, s, "Seguro de crédito acima do prémio mínimo? (1 = sim)", "", FMT_INT,
       lambda c, y: (f"IF(AND({c.l('ano', y)}>={c.p('sc_desde')},{c.p('sc_pct')}*({c.l('r_h', y)}+{c.l('r_d', y)})+{c.p('sc_pct_granel')}*{c.l('r_g', y)}"
                     f">{c.p('sc_min', y)}),1,0)"), style="memo")
    ln("wt_d", KF, s, "Prémio adicional pago pelos distribuidores/exportadores", "€/kg", FMT_EKG,
       lambda c, y: (f"IF({c.l('v_d_t', y)}>0,{c.l('wt_gap', y)}/({c.l('v_d_t', y)}*1000*(1-{c.p('d_desc')})*(1-{c.p('d_prov')}-{c.l('wt_ins', y)}*{c.p('sc_pct')})),\"n/a\")"))
    ln("wt_h", KF, s, "Referência HoReCa documentada necessária", "€/kg", FMT_EKG,
       lambda c, y: (f"IF({c.l('v_h_t', y)}>0,{c.p('h_ref')}+{c.l('wt_gap', y)}/({c.l('v_h_t', y)}*1000*(1-{c.p('h_desc')})*(1-{c.p('h_prov')}-{c.l('wt_ins', y)}*{c.p('sc_pct')})),\"n/a\")"))
    ln("wt_g_m", KF, s, "Margem do granel depois do seguro de crédito", "€/kg", FMT_EKG,
       lambda c, y: f"{c.l('g_margem', y)}-{c.l('wt_ins', y)}*{c.p('sc_pct_granel')}*{c.l('g_price', y)}")
    ln("wt_g", KF, s, "Volume de granel certificado necessário", "t", FMT_T,
       lambda c, y: f"IF({c.l('wt_g_m', y)}>0,{c.l('v_g_t', y)}+{c.l('wt_gap', y)}/({c.l('wt_g_m', y)}*1000),\"n/a\")")
    ln("wt_fix", KF, s, "Corte necessário nos custos fixos (pessoal + opex sem seguro de crédito)", "%", FMT_PCT,
       lambda c, y: f"{c.l('wt_gap', y)}/({c.l('pess_total', y)}+{c.l('ox_linhas', y)}+{c.l('ox_spv', y)})")
    ln("wt_tp_rate", KF, s, "Margem líquida da verificação para terceiros por € de carteira", "%", FMT_PCT2,
       lambda c, y: (f"{c.p('tp_taxa_util')}*{c.p('tp_rot')}+{c.p('tp_taxa_mon')}"
                     f"-{c.p('tp_share')}*{c.p('tp_custo_ver')}*{c.p('tp_rot')}/({c.p('tp_val_col')}*{c.p('tp_ltv')})"), style="memo")
    ln("wt_tp", KF, s, "Carteira de terceiros necessária", "€", FMT_EUR,
       lambda c, y: f"IF({c.l('wt_tp_rate', y)}>0,{c.l('v_tp_cart', y)}+{c.l('wt_gap', y)}/{c.l('wt_tp_rate', y)},\"n/a\")")
    return L


PRE_IDS = None  # preenchido em main() a partir do JSON


# ----------------------------------------------------------------------------- disposição
DETAIL_SHEETS = ["Volumes", "Receitas", "Pessoal", "Custos", "DR", "Cash flow", "Capital & financiamento"]
DETAIL_TITLES = {
    "Volumes": ("Volumes", "Volumes físicos por mercadoria e canal, clientes e carteira de terceiros (cenário ativo)."),
    "Receitas": ("Receitas", "Receitas por fluxo, GMV, capital de lotes financiado e margem de contribuição por fluxo (cenário ativo)."),
    "Pessoal": ("Pessoal", "Equipa por função; salário bruto anual com 14 meses, TSU de 23,75% + 2% e subsídio de refeição (cenário ativo)."),
    "Custos": ("Custos", "Economia por kg de cada canal, custos diretos, opex não salarial, capex e amortizações (cenário ativo)."),
    "DR": ("Demonstração de resultados", "Receita, margem bruta, margem de contribuição, EBITDA, amortizações, EBIT, IRC com derramas e resultado líquido (cenário ativo)."),
    "Cash flow": ("Cash flow", "Capital de giro por prazos médios, obrigações e float das séries, cash flow para o acionista, caixa e rondas (cenário ativo)."),
    "Capital & financiamento": ("Capital & financiamento", "Necessidade de capital, aplicação de fundos, runway, break-even, investidores e o que tem de ser verdade (cenário ativo)."),
}
DETAIL_START = {}


class Layout:
    """Mapa linha → (folha, linha) numa disposição."""

    def __init__(self):
        self.pos = {}


def layout_detail(lines):
    lay = Layout()
    rows = {}
    for sh in DETAIL_SHEETS:
        rows[sh] = DETAIL_START.get(sh, 5)
    cur_sec = {sh: None for sh in DETAIL_SHEETS}
    sec_rows = []
    for ln_ in lines:
        sh = ln_.sheet
        if ln_.section != cur_sec[sh]:
            if cur_sec[sh] is not None:
                rows[sh] += 1
            cur_sec[sh] = ln_.section
            sec_rows.append((sh, rows[sh], ln_.section))
            rows[sh] += 1
        lay.pos[ln_.lid] = (sh, rows[sh])
        rows[sh] += 1
    lay.sections = sec_rows
    return lay


def layout_motor_block(lines, start):
    lay = Layout()
    r = start
    cur = None
    secs = []
    for ln_ in lines:
        if ln_.section != cur:
            cur = ln_.section
            secs.append(("Motor", r, ln_.sheet + " · " + ln_.section))
            r += 1
        lay.pos[ln_.lid] = ("Motor", r)
        r += 1
    lay.sections = secs
    lay.end = r
    return lay


class Ctx:
    def __init__(self, preg, scen, layout, overrides=None):
        self.preg, self.scen, self.lay = preg, scen, layout
        self.ov = overrides or {}
        self.cur = None

    def p(self, pid, y=0):
        ref = self.preg.ref(pid, self.scen, y)
        if pid in self.ov:
            return "(" + self.ov[pid](ref, y) + ")"
        return ref

    def l(self, lid, y):
        sh, row = self.lay.pos[lid]
        ref = f"{YC[y]}{row}" if not LINES_BY[lid].scalar else f"C{row}"
        return ref if sh == self.cur else f"{q(sh)}!{ref}"

    def rng(self, lid):
        sh, row = self.lay.pos[lid]
        ref = f"$C${row}:$G${row}"
        return ref if sh == self.cur else f"{q(sh)}!{ref}"


LINES_BY = {}


def write_line_cells(ws, ctx, ln_, row):
    ctx.cur = ws.title
    put(ws, f"A{row}", ln_.label, bold=ln_.style == "key", italic=ln_.style == "memo",
        color=(MUTED if ln_.style in ("memo", "check") else None))
    put(ws, f"B{row}", ln_.unit, color=MUTED)
    cols = ["C"] if ln_.scalar else YC
    for y, col in enumerate(cols):
        v = ln_.fn(ctx, y)
        if isinstance(v, (int, float)) and ln_.style == "year":
            c = put(ws, f"{col}{row}", v, fmt=ln_.fmt, bold=True, color=MUTED)
        else:
            c = put(ws, f"{col}{row}", "=" + str(v), fmt=ln_.fmt, bold=ln_.style == "key")
        if ln_.style == "check":
            c.font = font(MUTED, italic=True)
        c.alignment = Alignment(horizontal="left" if ln_.scalar else "right")
    if ln_.style == "key":
        for col in ["A", "B"] + cols:
            ws[f"{col}{row}"].border = B_TOP
    if ln_.note:
        put(ws, f"H{row}", ln_.note, color=MUTED, italic=True)


def setup_detail_sheet(ws, sheet):
    t, sub = DETAIL_TITLES[sheet]
    title(ws, t, sub)
    put(ws, "A3", '="Cenário ativo: "&Pressupostos!$C$4&"  ·  euros correntes, sem IVA"', color=ACCENT, bold=True)
    header_row(ws, 4, ["Rubrica", "Unidade"] + [str(y) for y in YEARS] + ["Notas"])
    for k, v in {"A": 88, "B": 10, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 70}.items():
        ws.column_dimensions[k].width = v
    ws.freeze_panes = "C5"


def write_detail(wb, lines, preg, lay):
    sheets = {}
    for sh in DETAIL_SHEETS:
        ws = wb.create_sheet(sh)
        setup_detail_sheet(ws, sh)
        sheets[sh] = ws
    for sh, row, sec in lay.sections:
        section_row(sheets[sh], row, sec)
    ctx = Ctx(preg, 0, lay)
    for ln_ in lines:
        sh, row = lay.pos[ln_.lid]
        write_line_cells(sheets[sh], ctx, ln_, row)
    # Capital & financiamento: escalares numa coluna; rótulo de valor
    return sheets


def write_pessoal_table(ws, preg, P):
    r0 = ws.max_row + 2
    section_row(ws, r0, "Tabela salarial (inputs em Pressupostos; valores de 2027)")
    header_row(ws, r0 + 1, ["Função", "Fundador?", "Salário anual (14 meses)", "Salário mensal (÷ 14)", "Variável", "Custo anual c/ encargos (1 FTE)"])
    r = r0 + 2
    for fn, lab in ROLE_LABEL.items():
        put(ws, f"A{r}", lab)
        put(ws, f"B{r}", "sim (60% até ao gate)" if fn in FOUNDERS else "não", color=MUTED)
        put(ws, f"C{r}", "=" + preg.ref("sal_" + fn, 0), fmt=FMT_EUR)
        put(ws, f"D{r}", f"=C{r}/14", fmt=FMT_EUR)
        put(ws, f"E{r}", "=" + preg.ref("var_" + fn, 0), fmt=FMT_PCT)
        put(ws, f"F{r}", f"=C{r}*(1+E{r})*(1+{preg.ref('tsu', 0)}+{preg.ref('outros_enc', 0)})+{preg.ref('subs', 0)}", fmt=FMT_EUR)
        r += 1


# ----------------------------------------------------------------------------- casos do Motor
def motor_cases():
    """(chave, rótulo, cenário 1..4, overrides)."""
    cases = []
    for i, s in enumerate(SCEN):
        cases.append((s, SCEN_LONG[i], i + 1, {}))
    vol_ids = ["h_vol", "d_vol", "g_vol", "ag_pal", "cf_vol", "h_cli"]
    price_ids = ["az_indice", "h_ref", "d_ref"]
    for i in (2, 4):
        s = SCEN[i - 1]
        lab = SCEN_LONG[i - 1]
        for j in (1, 2, 3):
            cases.append((f"{s}__desc_{j}", f"{lab} · desconto HoReCa caso {j}", i,
                          {"h_desc": (lambda ref, y, j=j: f"Pressupostos!$C${PREG.by['sens_desc_' + str(j)].row}")}))
        for tag, pid in (("vol_neg", "sens_vol_neg"), ("vol_pos", "sens_vol_pos")):
            ov = {v: (lambda ref, y, pid=pid: f"{ref}*(1+Pressupostos!$C${PREG.by[pid].row})") for v in vol_ids}
            cases.append((f"{s}__{tag}", f"{lab} · volumes {'−' if tag == 'vol_neg' else '+'}", i, ov))
        ov = {v: (lambda ref, y: f"{ref}*(1+Pressupostos!$C${PREG.by['sens_preco'].row})") for v in price_ids}
        cases.append((f"{s}__preco", f"{lab} · preço do AOVE", i, ov))
        ov = {v: (lambda ref, y: f"{ref}*(1+Pressupostos!$C${PREG.by['sens_vol_pos'].row})") for v in vol_ids}
        ov["d_premio_venda"] = lambda ref, y: f"Pressupostos!$C${PREG.by['sens_premio_topo'].row}"
        ov["tp_cart"] = lambda ref, y: PREG.ref("tp_cart", 3, y)
        cases.append((f"{s}__topo", f"{lab} · topo dos intervalos (prémio 0,20; carteira do otimista; volumes +30%)", i, ov))
    return cases


def write_motor(wb, lines, preg):
    ws = wb.create_sheet("Motor")
    title(ws, "Motor de cálculo", "As mesmas linhas das folhas de detalhe, em blocos fixos por cenário e por caso de sensibilidade. "
                                  "Não editar: os inputs estão em Pressupostos. Cenários, Unit economics e Sensibilidade leem daqui.")
    for k, v in {"A": 88, "B": 10, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 18}.items():
        ws.column_dimensions[k].width = v
    header_row(ws, 4, ["Rubrica", "Unidade"] + [str(y) for y in YEARS] + ["id da linha"])
    ws.freeze_panes = "C5"
    r = 6
    blocks = {}
    for key, lab, scen, ov in motor_cases():
        c = ws.cell(row=r, column=1, value=f"Bloco: {lab}")
        c.font = Font(name=ARIAL, size=12, bold=True, color=INK)
        c.fill = FILL_PANEL
        for j in range(2, 9):
            ws.cell(row=r, column=j).fill = FILL_PANEL
        lay = layout_motor_block(lines, r + 1)
        for _, srow, sec in lay.sections:
            section_row(ws, srow, sec)
        ctx = Ctx(preg, scen, lay, ov)
        for ln_ in lines:
            _, row = lay.pos[ln_.lid]
            write_line_cells(ws, ctx, ln_, row)
            put(ws, f"H{row}", ln_.lid, color=MUTED, italic=True)
            ws.row_dimensions[row].outlineLevel = 1
        blocks[key] = lay
        r = lay.end + 2
    return ws, blocks


def mref(blocks, key, lid, y=None, absolute=True):
    sh, row = blocks[key].pos[lid]
    col = "C" if (y is None or LINES_BY[lid].scalar) else YC[y]
    return f"Motor!${col}${row}" if absolute else f"Motor!{col}{row}"


def mrange(blocks, key, lid):
    _, row = blocks[key].pos[lid]
    return f"Motor!$C${row}:$G${row}"


# ----------------------------------------------------------------------------- Unit economics
UE_POS = {}


def write_unit_economics(wb, preg, blocks):
    ws = wb.create_sheet("Unit economics", 2)
    title(ws, "Unit economics", "Lotes-tipo de ponta a ponta, por cenário (independente do seletor). Valores por kg lidos do Motor no ano de regime; "
                                "lote = €/kg × dimensão do lote.")
    for k, v in {"A": 80, "B": 9}.items():
        ws.column_dimensions[k].width = v
    for col in "CDEFGHIJ":
        ws.column_dimensions[col].width = 13
    ws.column_dimensions["K"].width = 60
    put(ws, "A3", "Ano de regime dos lotes-tipo", bold=True)
    put(ws, "C3", "=" + preg.ref("ue_ano", 0), fmt=FMT_YEAR)
    put(ws, "D3", "=C3-2026", fmt="0")
    put(ws, "E3", "← índice do ano (1 = 2027); mude o ano em Pressupostos", color=MUTED, italic=True)
    yidx = "$D$3"
    r = 5

    def at(key, lid, idx=yidx):
        return f"INDEX({mrange(blocks, key, lid)},1,{idx})"

    def ser_at(pid, s_i, idx=yidx):
        rr = preg.row_range(pid, s_i)
        return f"INDEX({rr},1,{idx})" if rr else preg.ref(pid, s_i)

    def lot_block(r, k, nm, lot_pid, extra_investor=True):
        section_row(ws, r, nm, ncols=11)
        r += 1
        hdr = ["Rubrica", "Unidade"]
        for lab in SCEN_LABEL:
            hdr += [f"{lab}: € / lote", f"{lab}: €/kg"]
        header_row(ws, r, hdr + ["Notas"])
        r += 1
        put(ws, f"A{r}", "Dimensão do lote")
        put(ws, f"B{r}", "kg", color=MUTED)
        kg_cells = {}
        for i, s in enumerate(SCEN):
            ce, ck = "CEGI"[i], "DFHJ"[i]
            put(ws, f"{ce}{r}", f"={preg.ref(lot_pid, 0)}*1000", fmt=FMT_INT)
            kg_cells[s] = f"${ce}${r}"
        kg_row = r
        r += 1
        rows = {}
        spec = [
            ("ref", "Preço de referência do comprador (mercado)", lambda s: at(s, f"{k}_ref"), 1, None),
            ("desc", "Desconto ao comprador", lambda s: f"-({at(s, f'{k}_ref')}-{at(s, f'{k}_price')})", 1, None),
            ("price", "Receita Sprout (preço de venda)", lambda s: at(s, f"{k}_price"), 1, "key"),
            ("gr", "Preço ao produtor (índice + prémio; pago ≤ 30 dias)", lambda s: f"-{at(s, f'{k}_gr')}", 1, None),
            ("emb", "Embalamento no lagar", lambda s: f"-{at(s, f'{k}_emb')}", 1, None),
            ("log", "Logística", lambda s: f"-{at(s, f'{k}_log')}", 1, None),
            ("3pl", "Armazém logístico (3PL) e picking", lambda s: f"-{at(s, f'{k}_3pl')}", 1, None),
            ("arm", "Armazenagem", lambda s: f"-{at(s, f'{k}_arm')}", 1, None),
            ("seg", "Seguro do stock", lambda s: f"-{at(s, f'{k}_seg')}", 1, None),
            ("anal", "Análises físico-químicas e painel (testes de qualidade)", lambda s: f"-{at(s, f'{k}_anal')}", 1, None),
            ("perdas", "Perdas físicas", lambda s: f"-{at(s, f'{k}_perdas')}", 1, None),
            ("descl", "Desclassificação esperada", lambda s: f"-{at(s, f'{k}_descl')}", 1, None),
            ("cust", "Custódia independente (collateral manager)", lambda s: f"-{at(s, f'{k}_cust')}", 1, None),
            ("pag", "Pagamentos bancários (SEPA; KYC/KYB no custo fixo)", lambda s: f"-{at(s, f'{k}_pag')}", 1, None),
            ("prov", "Provisão de crédito", lambda s: f"-{at(s, f'{k}_prov')}", 1, None),
            ("hedge", "Cobertura de preço explícita (azeite: sem futuros)", lambda s: f"-{at(s, f'{k}_hedge')}", 1, None),
            ("cupao", "Cupão das obrigações (retorno do investidor no ciclo)", lambda s: f"-{at(s, f'{k}_cupao')}", 1, None),
            ("margem", "Margem de contribuição do Sprout (spread da plataforma)", lambda s: at(s, f"{k}_margem"), 1, "key"),
        ]
        for key, lab, f, _, sty in spec:
            put(ws, f"A{r}", lab, bold=sty == "key")
            put(ws, f"B{r}", "€", color=MUTED)
            for i, s in enumerate(SCEN):
                ce, ck = "CEGI"[i], "DFHJ"[i]
                put(ws, f"{ck}{r}", "=" + f(s), fmt=FMT_EKG, bold=sty == "key")
                put(ws, f"{ce}{r}", f"={ck}{r}*{kg_cells[s]}", fmt=FMT_EUR, bold=sty == "key")
                if sty == "key":
                    ws[f"{ce}{r}"].border = B_TOP
                    ws[f"{ck}{r}"].border = B_TOP
            rows[key] = r
            r += 1
        put(ws, f"A{r}", "Margem em % da receita")
        put(ws, f"B{r}", "%", color=MUTED)
        for i in range(4):
            ck = "DFHJ"[i]
            put(ws, f"{ck}{r}", f"=IF({ck}{rows['price']}>0,{ck}{rows['margem']}/{ck}{rows['price']},0)", fmt=FMT_PCT)
        rows["margem_pct"] = r
        r += 1
        # investidor e reserva
        put(ws, f"A{r}", "Lado do investidor e almofada", bold=True, color=MUTED)
        r += 1
        inv = [
            ("cap", "Capital imobilizado no lote (granel + embalamento)", "€",
             lambda s, si: f"({at(s, f'{k}_gr')}+{at(s, f'{k}_emb')})*{kg_cells[s]}", FMT_EUR),
            ("ciclo", "Ciclo líquido do capital (stock + recebimento − pagamento)", "dias",
             lambda s, si: f"{at(s, f'{k}_ds')}+{at(s, f'{k}_dr')}-{preg.ref('az_dp', 0)}", FMT_DAYS),
            ("pct", "Parte do capital de giro financiada por obrigações", "%", lambda s, si: ser_at("fin_pct", si), FMT_PCT),
            ("notes", "Obrigações afetas ao lote (capital levantado = financiado ÷ utilização)", "€",
             lambda s, si: f"{ser_at('fin_pct', si)}*({at(s, f'{k}_gr')}+{at(s, f'{k}_emb')})*{kg_cells[s]}/{preg.ref('fin_util', si)}", FMT_EUR),
            ("sub", "Capital próprio do patrocinador no lote (subordinado)", "€",
             lambda s, si: f"(1-{ser_at('fin_pct', si)})*({at(s, f'{k}_gr')}+{at(s, f'{k}_emb')})*{kg_cells[s]}", FMT_EUR),
            ("cup", "Cupão bruto anual", "%", lambda s, si: ser_at("fin_cupao", si), FMT_PCT2),
            ("cupliq", "Retorno líquido anual de um particular residente (retenção de 28%)", "%",
             lambda s, si: f"{ser_at('fin_cupao', si)}*(1-{preg.ref('fin_ret_sing', 0)})", FMT_PCT2),
            ("cupciclo", "Cupão pago no ciclo do lote", "€", lambda s, si: f"{at(s, f'{k}_cupao')}*{kg_cells[s]}", FMT_EUR),
            ("retciclo", "Retorno bruto do investidor no ciclo (% das obrigações afetas ao lote)", "%",
             lambda s, si: (f"IF({ser_at('fin_pct', si)}>0,{at(s, f'{k}_cupao')}*{kg_cells[s]}/({ser_at('fin_pct', si)}*({at(s, f'{k}_gr')}+{at(s, f'{k}_emb')})"
                            f"*{kg_cells[s]}/{preg.ref('fin_util', si)}),0)"), FMT_PCT2),
            ("res", "Retida na reserva de primeira perda (30% da margem; teto 5% das obrigações do lote)", "€",
             lambda s, si: (f"MAX(0,MIN({preg.ref('fin_res_pct_margem', 0)}*{at(s, f'{k}_margem')}*{kg_cells[s]},"
                            f"{preg.ref('fin_res_pct_notes', 0)}*{ser_at('fin_pct', si)}*({at(s, f'{k}_gr')}+{at(s, f'{k}_emb')})*{kg_cells[s]}/{preg.ref('fin_util', si)}))"), FMT_EUR),
            ("livre", "Margem livre para a OpCo depois da reserva", "€",
             lambda s, si: f"{at(s, f'{k}_margem')}*{kg_cells[s]}-RES", FMT_EUR),
            ("kyc", "Memo: KYC/KYB e screening das contrapartes do lote (no custo fixo de tecnologia)", "€",
             (lambda s, si: f"{kg_cells[s]}/({preg.ref('az_cons_cliente', 0)}*1000)*({preg.ref('kyc_kyb', 0)}+12*{preg.ref('kyc_screen', 0)})") if k == "h"
             else (lambda s, si: f"{preg.ref('kyc_kyb', 0)}+12*{preg.ref('kyc_screen', 0)}"), FMT_EUR),
            ("noc", "Memo: custo nocional do capital próprio (8%/ano; fora do EBITDA)", "€",
             lambda s, si: (f"{preg.ref('fin_nocional', 0)}*(1-{ser_at('fin_pct', si)})*({at(s, f'{k}_gr')}+{at(s, f'{k}_emb')})"
                            f"*({at(s, f'{k}_ds')}+{at(s, f'{k}_dr')}-{preg.ref('az_dp', 0)})/365*{kg_cells[s]}"), FMT_EUR),
        ]
        for key, lab, unit, f, fmt in inv:
            put(ws, f"A{r}", lab)
            put(ws, f"B{r}", unit, color=MUTED)
            for i, s in enumerate(SCEN):
                ce = "CEGI"[i]
                fx = f(s, i + 1)
                if key == "livre":
                    fx = fx.replace("RES", f"{ce}{rows['res']}")
                put(ws, f"{ce}{r}", "=" + fx, fmt=fmt)
            rows[key] = r
            r += 1
        return r + 1, rows

    # A. lote HoReCa
    put(ws, f"A{r}", "A. Lote de azeite HoReCa/gourmet (origem identificada; 5 L PET/BIB e vidro)", bold=True, color=INK)
    r += 1
    r, rows_h = lot_block(r, "h", "Ponta a ponta no ano de regime (por lote e por kg)", "h_lote")
    UE_POS["h"] = rows_h
    # bolo de intermediação
    section_row(ws, r, "Bolo de intermediação (Π = referência − granel − embalamento − logística) e repartição", ncols=11)
    r += 1
    header_row(ws, r, ["Rubrica", "Unidade"] + [f"{lab}" for lab in SCEN_LABEL for _ in (0,)] + ["", "", "", "", "Notas"])
    r += 1
    pi = {}
    items = [
        ("pi", "Π por kg", "€/kg", lambda ck: f"={ck}{rows_h['ref']}+{ck}{rows_h['gr']}+{ck}{rows_h['emb']}+{ck}{rows_h['log']}", FMT_EKG),
        ("pi_desc", "Comprador: desconto", "% de Π", lambda ck: f"=IF(PI>0,-{ck}{rows_h['desc']}/PI,0)", FMT_PCT),
        ("pi_cup", "Investidores: cupão", "% de Π", lambda ck: f"=IF(PI>0,-{ck}{rows_h['cupao']}/PI,0)", FMT_PCT),
        ("pi_cd", "Custos diretos (3PL, custódia, qualidade, crédito, pagamentos, armazenagem)", "% de Π",
         lambda ck: ("=IF(PI>0,-(" + "+".join(f"{ck}{rows_h[x]}" for x in ("3pl", "arm", "seg", "anal", "perdas", "descl", "cust", "pag", "prov", "hedge")) + ")/PI,0)"), FMT_PCT),
        ("pi_sp", "Sprout: margem de contribuição", "% de Π", lambda ck: f"=IF(PI>0,{ck}{rows_h['margem']}/PI,0)", FMT_PCT),
    ]
    for key, lab, unit, f, fmt in items:
        put(ws, f"A{r}", lab)
        put(ws, f"B{r}", unit, color=MUTED)
        for i in range(4):
            ck = "DFHJ"[i]
            col = "CDEF"[i]
            fx = f(ck).replace("PI", f"{col}{pi['pi']}" if "pi" in pi else "0")
            put(ws, f"{col}{r}", fx, fmt=fmt)
        pi[key] = r
        r += 1
    UE_POS["pi"] = pi
    r += 1
    # tese literal
    section_row(ws, r, "Teste da tese literal do documento-fonte (lote HoReCa, por kg): captura de 80–90% da margem do embalador + desconto de 10–15%", ncols=11)
    r += 1
    header_row(ws, r, ["Rubrica", "Unidade"] + SCEN_LABEL + ["", "", "", "", "Notas"])
    r += 1
    tese = {}
    titems = [
        ("m", "Margem líquida do embalador (média de 3 campanhas do Observatorio)", lambda i, ck: f"={preg.ref('tese_m_med', 0)}"),
        ("cap", "Captura de 85% (ponto médio de 80–90%)", lambda i, ck: f"={preg.ref('tese_cap_2', 0)}*{'CDEF'[i]}{{m}}"),
        ("desc", "Custo do desconto ao comprador (desconto × referência)", lambda i, ck: f"={ck}{rows_h['desc']}"),
        ("cup", "Retorno do investidor (cupão do ciclo)", lambda i, ck: f"={ck}{rows_h['cupao']}"),
        ("cd", "Custos diretos da plataforma (sem granel, embalamento e logística)", lambda i, ck: (
            "=" + "+".join(f"{ck}{rows_h[x]}" for x in ("3pl", "arm", "seg", "anal", "perdas", "descl", "cust", "pag", "prov", "hedge")))),
        ("spread", "Spread que sobra para a plataforma", lambda i, ck: f"={'CDEF'[i]}{{cap}}+{'CDEF'[i]}{{desc}}+{'CDEF'[i]}{{cup}}+{'CDEF'[i]}{{cd}}"),
        ("ratio", "Desconto ÷ captura (vezes)", lambda i, ck: f"=IF({'CDEF'[i]}{{cap}}>0,-{'CDEF'[i]}{{desc}}/{'CDEF'[i]}{{cap}},0)"),
        ("verd", "Veredicto", lambda i, ck: f"=IF({'CDEF'[i]}{{spread}}<0,\"não fecha\",\"fecha\")"),
    ]
    for key, lab, f in titems:
        put(ws, f"A{r}", lab, bold=key in ("spread", "verd"))
        put(ws, f"B{r}", "x" if key == "ratio" else ("" if key == "verd" else "€/kg"), color=MUTED)
        tese[key] = r
        for i in range(4):
            ck = "DFHJ"[i]
            fx = f(i, ck).format(**tese)
            put(ws, f"{'CDEF'[i]}{r}", fx, fmt=(FMT_NUM if key == "ratio" else FMT_EKG), bold=key in ("spread", "verd"))
        r += 1
    UE_POS["tese"] = tese
    put(ws, f"K{tese['spread']}", "Negativo = a tese literal não fecha: o desconto custa mais do que a margem capturada", color=MUTED, italic=True)
    r += 1
    # risco de preço
    section_row(ws, r, "Risco de preço do lote HoReCa (por que só back-to-back)", ncols=11)
    r += 1
    header_row(ws, r, ["Rubrica", "Unidade"] + SCEN_LABEL + ["", "", "", "", "Notas"])
    r += 1
    rk = {}
    ritems = [
        ("val", "Valor do stock ao custo (granel + embalamento)", "€", lambda i: f"={'CEGI'[i]}{rows_h['cap']}", FMT_EUR),
        ("s1_100", "Choque de −30% no preço: perda com 100% do lote em aberto", "€", lambda i: f"={preg.ref('sens_preco', 0)}*{'CDEF'[i]}{{val}}", FMT_EUR),
        ("s1_15", "Idem com 15% em aberto (política back-to-back)", "€", lambda i: f"={preg.ref('exp_aberta', 0)}*{'CDEF'[i]}{{s1_100}}", FMT_EUR),
        ("s1_marg", "Perda com 100% em aberto ÷ margem do lote", "x", lambda i: f"=IF({'CEGI'[i]}{rows_h['margem']}>0,-{'CDEF'[i]}{{s1_100}}/{'CEGI'[i]}{rows_h['margem']},0)", FMT_NUM),
        ("h18_lo", "Detenção de 18 meses sem cobertura: carry + VaR95 (σ 23%) − margem", "% do lote",
         lambda i: f"={preg.ref('s6_carry', 0)}+{preg.ref('s6_var_lo', 0)}-{'DFHJ'[i]}{rows_h['margem_pct']}", FMT_PCT),
        ("h18_hi", "Idem com σ 42%", "% do lote", lambda i: f"={preg.ref('s6_carry', 0)}+{preg.ref('s6_var_hi', 0)}-{'DFHJ'[i]}{rows_h['margem_pct']}", FMT_PCT),
    ]
    for key, lab, unit, f, fmt in ritems:
        put(ws, f"A{r}", lab)
        put(ws, f"B{r}", unit, color=MUTED)
        rk[key] = r
        for i in range(4):
            put(ws, f"{'CDEF'[i]}{r}", f(i).format(**rk), fmt=fmt)
        r += 1
    UE_POS["risco"] = rk
    put(ws, f"K{rk['h18_lo']}", "40 §5.2 e §1.3; 01 §2.6. Não há futuros de azeite para cobrir esta perda.", color=MUTED, italic=True)
    r += 1
    # piloto
    section_row(ws, r, "Piloto de 2027 (sem obrigações: stock em capital próprio)", ncols=11)
    r += 1
    header_row(ws, r, ["Rubrica", "Unidade"] + SCEN_LABEL + ["", "", "", "", "Notas"])
    r += 1
    pil = {}
    for k, nm in CH[:2]:
        for key, lab, lid, fmt in (("price", "preço de venda", f"{k}_price", FMT_EKG), ("custo", "custo direto", f"{k}_custo", FMT_EKG),
                                   ("margem", "margem de contribuição", f"{k}_margem", FMT_EKG), ("pct", "margem em %", f"{k}_margem_pct", FMT_PCT)):
            put(ws, f"A{r}", f"{nm[0].upper() + nm[1:]}: {lab}")
            put(ws, f"B{r}", "%" if key == "pct" else "€/kg", color=MUTED)
            for i, s in enumerate(SCEN):
                put(ws, f"{'CDEF'[i]}{r}", "=" + mref(blocks, s, lid, 0), fmt=fmt)
            pil[f"{k}_{key}"] = r
            r += 1
    UE_POS["piloto"] = pil
    r += 1
    # B. distribuidores, C. granel
    put(ws, f"A{r}", "B. Lote de azeite para distribuidores/exportação (paletes de 5 L; preço de saída do embalador)", bold=True, color=INK)
    r += 1
    r, rows_d = lot_block(r, "d", "Ponta a ponta no ano de regime (por lote e por kg)", "d_lote")
    UE_POS["d"] = rows_d
    put(ws, f"A{r}", "C. Lote de granel certificado (IBC/cisterna; só o Recomendado tem volume)", bold=True, color=INK)
    r += 1
    r, rows_g = lot_block(r, "g", "Ponta a ponta no ano de regime (por lote e por kg)", "g_lote")
    UE_POS["g"] = rows_g
    # D. água
    put(ws, f"A{r}", "D. Lote de água: 20 paletes (comparação principal vs. agente; recomendado: agente)", bold=True, color=INK)
    r += 1
    header_row(ws, r, ["Rubrica", "Unidade", "Principal, desconto 10%", "Principal, desconto 15%"] + [f"Agente: {x}" for x in SCEN_LABEL] + ["", "Notas"])
    r += 1
    ag = {}
    L_ = f"{preg.ref('ag_lote_pal', 0)}*{preg.ref('ag_l_pal', 0)}"
    base_cup = f"INDEX({preg.row_range('fin_cupao', 2)},1,1)"
    rows_spec = [
        ("litros", "Volume", "L", [f"={L_}"] * 2 + [f"={L_}"] * 4, FMT_INT),
        ("val", "Valor à referência (preço B2B, NV)", "€", [f"=C{{litros}}*{preg.ref('ag_ref_l', 0)}", f"=D{{litros}}*{preg.ref('ag_ref_l', 0)}"]
         + [f"={col}{{litros}}*{preg.ref('ag_ref_l', 0)}" for col in "EFGH"], FMT_EUR),
        ("gmv", "Receita (principal) ou GMV (agente)", "€", [f"=C{{val}}*(1-{preg.ref('ag_desc', 0)})", f"=D{{val}}*(1-{preg.ref('ag_desc_alt', 0)})"]
         + [f"={col}{{val}}*(1-{preg.ref('ag_desc', 0)})" for col in "EFGH"], FMT_EUR),
        ("exfab", "Preço ex-fábrica", "€", [f"=-C{{litros}}*{preg.ref('ag_exfab', 0)}", f"=-D{{litros}}*{preg.ref('ag_exfab', 0)}"] + [""] * 4, FMT_EUR),
        ("log", "Logística (troncal + entrega)", "€", [f"=-({preg.ref('ag_troncal', 0)}+{preg.ref('ag_entrega', 0)}*C{{litros}})",
                                                       f"=-({preg.ref('ag_troncal', 0)}+{preg.ref('ag_entrega', 0)}*D{{litros}})"] + [""] * 4, FMT_EUR),
        ("ops", "Operação e pagamentos", "€", [f"=-{preg.ref('ag_ops', 0)}"] * 2 + [""] * 4, FMT_EUR),
        ("sdr", "Depósito SDR a adiantar e reconciliar", "€", [f"={preg.ref('ag_lote_pal', 0)}*{preg.ref('ag_garrafas', 0)}*{preg.ref('ag_sdr', 0)}"] * 2 + ["0"] * 4, FMT_EUR),
        ("fin", "Financiamento do stock e do SDR (cupão de 2027, 30 dias)", "€",
         [f"=-{base_cup}*(-C{{exfab}}+C{{sdr}})*{preg.ref('ag_dias_fin', 0)}/365", f"=-{base_cup}*(-D{{exfab}}+D{{sdr}})*{preg.ref('ag_dias_fin', 0)}/365"] + [""] * 4, FMT_EUR),
        ("prov", "Provisão de crédito (0,6%)", "€", [f"=-{preg.ref('h_prov', 2)}*C{{gmv}}", f"=-{preg.ref('h_prov', 2)}*D{{gmv}}"] + [""] * 4, FMT_EUR),
        ("marg", "Margem (principal) ou comissão (agente) do Sprout", "€",
         ["=C{gmv}+C{exfab}+C{log}+C{ops}+C{fin}+C{prov}", "=D{gmv}+D{exfab}+D{log}+D{ops}+D{fin}+D{prov}"]
         + [f"={col}{{gmv}}*{preg.ref('ag_com', i + 1)}" for i, col in enumerate("EFGH")], FMT_EUR),
        ("pct", "em % da receita/GMV", "%", [f"=IF({col}{{gmv}}>0,{col}{{marg}}/{col}{{gmv}},0)" for col in "CDEFGH"], FMT_PCT),
        ("pal", "por palete", "€/palete", [f"={col}{{marg}}/{preg.ref('ag_lote_pal', 0)}" for col in "CDEFGH"], FMT_EKG2),
        ("emb", "Margem de embalador capturável (embalamento na origem obrigatório, Dir. 2009/54/CE)", "€", ["0"] * 6, FMT_EUR),
    ]
    for key, lab, unit, fs, fmt in rows_spec:
        ag[key] = r
        put(ws, f"A{r}", lab, bold=key == "marg")
        put(ws, f"B{r}", unit, color=MUTED)
        for col, fx in zip("CDEFGH", fs):
            if fx == "":
                put(ws, f"{col}{r}", "—", color=MUTED, align=Alignment(horizontal="right"))
                continue
            fx = fx.format(**ag)
            if not fx.startswith("="):
                fx = "=" + fx
            put(ws, f"{col}{r}", fx, fmt=fmt, bold=key == "marg")
        r += 1
    UE_POS["agua"] = ag
    put(ws, f"K{ag['fin']}", "01 §2.7 indica 38 € (≈ 10%/ano); com o cupão de 8% dá ~30 €: ver Fontes", color=MUTED, italic=True)
    ws.freeze_panes = "C5"
    return ws


# ----------------------------------------------------------------------------- Cenários
CEN_POS = {}

CEN_METRICS = [
    ("v_az_t", "Azeite transacionado", "t", FMT_T),
    ("gmv_total", "GMV total", "€", FMT_EUR),
    ("r_total", "Receita total", "€", FMT_EUR),
    ("r_h", "— vendas de azeite HoReCa/gourmet", "€", FMT_EUR),
    ("r_d", "— vendas de azeite a distribuidores/exportação", "€", FMT_EUR),
    ("r_g", "— vendas de granel certificado", "€", FMT_EUR),
    ("r_cf", "— vendas de café verde", "€", FMT_EUR),
    ("r_ag", "— comissões de água", "€", FMT_EUR),
    ("r_tp", "— verificação para terceiros", "€", FMT_EUR),
    ("r_dd", "— dados e índice", "€", FMT_EUR),
    ("dr_mb", "Margem bruta", "€", FMT_EUR),
    ("mc_total", "Margem de contribuição", "€", FMT_EUR),
    ("mc_az_kg", "Margem média do azeite", "€/kg", FMT_EKG),
    ("pess_total", "Pessoal", "€", FMT_EUR),
    ("ox_total", "Opex não salarial", "€", FMT_EUR),
    ("dr_ebitda", "EBITDA", "€", FMT_EUR),
    ("dr_rl", "Resultado líquido", "€", FMT_EUR),
    ("cf_acion", "Cash flow para o acionista", "€", FMT_EUR),
    ("cf_acum", "Cash flow acumulado", "€", FMT_EUR),
    ("cx_plug", "Rondas adicionais necessárias (além da pré-seed)", "€", FMT_EUR),
    ("cx_fim", "Caixa livre no fim do ano", "€", FMT_EUR),
    ("ob", "Obrigações em circulação", "€", FMT_EUR),
    ("lotes_fin_obr", "Valor de lotes financiado por obrigações (capital financiado)", "€", FMT_EUR),
    ("inv_n", "Investidores estimados", "n.º", FMT_INT),
    ("fte_total", "FTE", "FTE", FMT_FTE),
]
CEN_SCALARS = [
    ("k_need", "Necessidade de capital próprio (pico no horizonte)", "€", FMT_EUR),
    ("k_min_ano", "Ano do mínimo do cash flow acumulado", "ano", FMT_YEAR),
    ("k_ainda", "O acumulado ainda desce em 2031?", "", None),
    ("k_rondas", "Capital próprio levantado (pré-seed + rondas adicionais)", "€", FMT_EUR),
    ("be_data", "Break-even do EBITDA (mês, run-rate)", "data", FMT_DATE),
    ("be_ano_ebitda", "1.º ano com EBITDA ≥ 0", "ano", FMT_YEAR),
    ("be_ano_cf", "1.º ano com cash flow para o acionista ≥ 0", "ano", FMT_YEAR),
    ("b_runway", "Runway da pré-seed (meses desde jan. 2027)", "meses", FMT_MES),
    ("b_gate", "Gasto até ao gate de 31-08-2027", "€", FMT_EUR),
]
CEN_WT = [
    ("wt_gap", "Défice de EBITDA a cobrir em 2031", "€", FMT_EUR),
    ("wt_d", "… ou prémio adicional dos distribuidores/exportadores", "€/kg", FMT_EKG),
    ("wt_h", "… ou referência HoReCa documentada", "€/kg", FMT_EKG),
    ("wt_g", "… ou volume de granel certificado", "t", FMT_T),
    ("wt_fix", "… ou corte nos custos fixos", "%", FMT_PCT),
    ("wt_tp", "… ou carteira de verificação para terceiros", "€", FMT_EUR),
]


def write_cenarios(wb, blocks, P, detail_lay):
    ws = wb.create_sheet("Cenários")
    title(ws, "Cenários", "Os três cenários do documento-fonte (com as correções verificadas, sem ajustes) e o cenário Recomendado (reestruturado), lado a lado. "
                          "Ligações ao Motor: não dependem do seletor.")
    for k, v in {"A": 62, "B": 26, "C": 20, "D": 20, "E": 20, "F": 20, "G": 20, "H": 100}.items():
        ws.column_dimensions[k].width = v
    r = 4
    section_row(ws, r, "1. Indicadores por ano")
    r += 1
    header_row(ws, r, ["Indicador", "Cenário"] + [str(y) for y in YEARS] + ["Unidade"])
    r += 1
    for lid, lab, unit, fmt in CEN_METRICS:
        CEN_POS[lid] = {}
        for i, s in enumerate(SCEN):
            put(ws, f"A{r}", lab if i == 0 else "", bold=i == 0)
            put(ws, f"B{r}", SCEN_LABEL[i], color=MUTED if i < 3 else ACCENT, bold=i == 3)
            for y in range(NY):
                put(ws, f"{YC[y]}{r}", "=" + mref(blocks, s, lid, y, absolute=False), fmt=fmt, bold=i == 3)
            put(ws, f"H{r}", unit, color=MUTED)
            CEN_POS[lid][s] = r
            r += 1
        for col in "ABCDEFGH":
            ws[f"{col}{r - 1}"].border = B_BOTTOM
    r += 1
    section_row(ws, r, "2. Capital, break-even e runway")
    r += 1
    header_row(ws, r, ["Indicador", "Unidade"] + SCEN_LABEL + ["", "Notas"])
    r += 1
    for lid, lab, unit, fmt in CEN_SCALARS:
        CEN_POS[lid] = r
        put(ws, f"A{r}", lab, bold=lid in ("k_need", "be_data"))
        put(ws, f"B{r}", unit, color=MUTED)
        for i, s in enumerate(SCEN):
            put(ws, f"{'CDEF'[i]}{r}", "=" + mref(blocks, s, lid, absolute=False), fmt=fmt, bold=lid in ("k_need", "be_data"),
                align=Alignment(horizontal="right"))
        r += 1
    put(ws, f"H{CEN_POS['be_data']}", "Run-rate mensal interpolado entre os pontos médios dos anos (estimativa)", color=MUTED, italic=True)
    put(ws, f"H{CEN_POS['k_need']}", "Pré-arranque + capital da SPV + caixa mínima + vale do cash flow acumulado", color=MUTED, italic=True)
    r += 1
    section_row(ws, r, "3. O que tem de ser verdade para EBITDA = 0 em 2031 (uma alavanca de cada vez, o resto igual)")
    r += 1
    header_row(ws, r, ["Alavanca", "Unidade"] + SCEN_LABEL + ["", "Leitura"])
    r += 1
    for lid, lab, unit, fmt in CEN_WT:
        CEN_POS[lid] = r
        put(ws, f"A{r}", lab)
        put(ws, f"B{r}", unit, color=MUTED)
        for i, s in enumerate(SCEN):
            put(ws, f"{'CDEF'[i]}{r}", "=" + mref(blocks, s, lid, 4, absolute=False), fmt=fmt, align=Alignment(horizontal="right"))
        r += 1
    put(ws, f"H{CEN_POS['wt_d']}", "Valor a somar ao prémio já assumido; comparar com 0,10–0,20 (condição do 01 para a seed)", color=MUTED, italic=True)
    put(ws, f"H{CEN_POS['wt_h']}", "Comparar com 5,3–6,0 €/kg (equivalente de retalho a out. 2026, NV)", color=MUTED, italic=True)
    put(ws, f"H{CEN_POS['wt_g']}", "Comparar com ~175 kt de produção portuguesa", color=MUTED, italic=True)
    r += 1
    section_row(ws, r, "4. O que muda no cenário Recomendado (reestruturado) face ao base")
    r += 1
    header_row(ws, r, ["Parâmetro", "Base", "Recomendado", "", "", "", "", "Justificação e fonte"])
    r += 1
    for a in P["cenario_recomendado"]["alteracoes"]:
        put(ws, f"A{r}", a["parametro"], align=WRAP)
        put(ws, f"B{r}", fmt_val(a["base"]), color=MUTED, align=WRAP)
        put(ws, f"C{r}", fmt_val(a["recomendado"]), color=ACCENT, align=WRAP)
        ws.merge_cells(f"C{r}:G{r}")
        put(ws, f"H{r}", a["justificacao"] + " Fonte: " + a["fonte"], color=MUTED, align=WRAP)
        ws.row_dimensions[r].height = 13 * max(2, (len(a["justificacao"]) + len(a["fonte"]) + 8) // 95 + 1)
        r += 1
    for lab, txt in (("Não alterado", P["cenario_recomendado"]["nao_alterado"]), ("Produto para investidores", P["cenario_recomendado"]["produto_investidor"])):
        put(ws, f"A{r}", lab, bold=True, align=Alignment(vertical="top"))
        put(ws, f"H{r}", txt, color=MUTED, align=WRAP)
        ws.row_dimensions[r].height = 13 * max(2, len(txt) // 95 + 1)
        r += 1
    r -= 1
    r += 2
    section_row(ws, r, "5. Verificações (devem ser 0)")
    r += 1
    sel = "Pressupostos!$E$4"
    checks = []
    for lid, lab in (("dr_ebitda", "EBITDA: folha DR − Motor (cenário ativo)"), ("cf_acum", "Cash flow acumulado: folha Cash flow − Motor")):
        put(ws, f"A{r}", lab)
        for y in range(NY):
            sh, row = detail_lay.pos[lid]
            det = f"{q(sh)}!{YC[y]}{row}"
            ch = "CHOOSE(" + sel + "," + ",".join(mref(blocks, s, lid, y) for s in SCEN) + ")"
            put(ws, f"{YC[y]}{r}", f"=ROUND({det}-{ch},4)", fmt=FMT_EUR)
        checks.append(r)
        r += 1
    sh, row = detail_lay.pos["k_need"]
    put(ws, f"A{r}", "Necessidade de capital: folha Capital − Motor")
    put(ws, f"C{r}", f"=ROUND({q(sh)}!C{row}-CHOOSE({sel}," + ",".join(mref(blocks, s, 'k_need') for s in SCEN) + "),4)", fmt=FMT_EUR)
    checks.append(r)
    r += 1
    put(ws, f"A{r}", "Verificações internas do Motor (margem e cash flow), soma dos valores absolutos")
    terms = []
    for key in blocks:
        for lid in ("chk_mc", "chk_cf"):
            terms.append(f"SUMPRODUCT(ABS({mrange(blocks, key, lid)}))")
    put(ws, f"C{r}", "=ROUND(" + "+".join(terms) + ",4)", fmt=FMT_EUR)
    checks.append(r)
    r += 1
    put(ws, f"A{r}", "Estado", bold=True)
    put(ws, f"C{r}", "=IF(" + "+".join(f"SUMPRODUCT(ABS(C{x}:G{x}))" for x in checks) + "=0,\"OK\",\"ERRO\")", bold=True)
    CEN_POS["estado"] = r
    ws.freeze_panes = "C6"
    return ws


def fmt_val(v):
    if isinstance(v, list):
        return " / ".join(fmt_num(x) for x in v)
    if isinstance(v, (int, float)):
        return fmt_num(v)
    return str(v)


def fmt_num(x, dec=None):
    """Formatação pt-PT: espaço fino? Não: ponto de milhar e vírgula decimal."""
    if isinstance(x, int) or (isinstance(x, float) and x.is_integer() and abs(x) >= 100):
        s = f"{int(round(x)):,}".replace(",", ".")
        return s
    if dec is None:
        dec = 2 if abs(x) < 100 else 0
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


# ----------------------------------------------------------------------------- Sensibilidade
SENS_POS = {}


def write_sensibilidade(wb, preg, blocks):
    ws = wb.create_sheet("Sensibilidade")
    title(ws, "Sensibilidade", "Base e Recomendado: desconto HoReCa, volumes e preço do AOVE (motor completo, recalculado por caso); "
                               "tese literal do documento-fonte: captura da margem do embalador × desconto.")
    for k, v in {"A": 80, "B": 14, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15, "I": 20}.items():
        ws.column_dimensions[k].width = v
    r = 4
    section_row(ws, r, "Casos (inputs em Pressupostos)", ncols=9)
    r += 1
    for pid in ("sens_desc_1", "sens_desc_2", "sens_desc_3", "sens_vol_neg", "sens_vol_pos", "sens_preco", "sens_premio_topo", "exp_aberta"):
        put(ws, f"A{r}", PREG.by[pid].label)
        put(ws, f"B{r}", "=" + preg.ref(pid, 0), fmt=(FMT_EKG2 if pid == "sens_premio_topo" else FMT_PCT))
        r += 1
    r += 1
    cols = [("mh", "Margem HoReCa 2029 (€/kg)", FMT_EKG), ("e29", "EBITDA 2029 (€)", FMT_EUR), ("e31", "EBITDA 2031 (€)", FMT_EUR),
            ("d31", "Δ EBITDA 2031 vs. cenário (€)", FMT_EUR), ("acum", "Cash flow acumulado 2031 (€)", FMT_EUR),
            ("need", "Necessidade de capital (€)", FMT_EUR), ("be", "Break-even do EBITDA", FMT_DATE)]
    for s in ("base", "recomendado"):
        lab = SCEN_LONG[SCEN.index(s)]
        section_row(ws, r, f"Cenário {lab}", ncols=9)
        r += 1
        header_row(ws, r, ["Caso", "Valor do caso"] + [c[1] for c in cols])
        r += 1
        SENS_POS[s] = {}
        cases = [(s, "Cenário (referência)", None), (f"{s}__desc_1", "Desconto HoReCa", "sens_desc_1"),
                 (f"{s}__desc_2", "Desconto HoReCa", "sens_desc_2"), (f"{s}__desc_3", "Desconto HoReCa", "sens_desc_3"),
                 (f"{s}__vol_neg", "Volumes físicos (estrutura fixa)", "sens_vol_neg"), (f"{s}__vol_pos", "Volumes físicos (estrutura fixa)", "sens_vol_pos"),
                 (f"{s}__preco", "Preço do AOVE (índice e referências; repasse proporcional)", "sens_preco"),
                 (f"{s}__topo", "Topo dos intervalos: prémio de origem dos distribuidores, carteira de terceiros do otimista, volumes +30%", "sens_premio_topo")]
        ref_row = r
        for key, clab, pid in cases:
            put(ws, f"A{r}", clab, bold=key == s)
            if pid:
                put(ws, f"B{r}", "=" + preg.ref(pid, 0), fmt=(FMT_EKG2 if pid == "sens_premio_topo" else FMT_PCT), align=Alignment(horizontal="right"))
            vals = {
                "mh": "=" + mref(blocks, key, "h_margem", 2),
                "e29": "=" + mref(blocks, key, "dr_ebitda", 2),
                "e31": "=" + mref(blocks, key, "dr_ebitda", 4),
                "d31": f"=E{r}-E{ref_row}",
                "acum": "=" + mref(blocks, key, "cf_acum", 4),
                "need": "=" + mref(blocks, key, "k_need"),
                "be": "=" + mref(blocks, key, "be_data"),
            }
            for j, (ck, _, fmt) in enumerate(cols):
                put(ws, f"{'CDEFGHI'[j]}{r}", vals[ck], fmt=fmt, bold=key == s, align=Alignment(horizontal="right"))
            SENS_POS[s][key] = r
            r += 1
        r += 1
    # tese literal
    section_row(ws, r, "Tese literal do documento-fonte, lote HoReCa do base em 2029 (€/kg que sobram para a plataforma)", ncols=9)
    r += 1
    put(ws, f"A{r}", "Saldo = captura × margem do embalador − desconto × referência − cupão − custos diretos da plataforma", color=MUTED, italic=True)
    r += 1
    ref = mref(blocks, "base", "h_ref", 2)
    cup = mref(blocks, "base", "h_cupao", 2)
    cd = "+".join(mref(blocks, "base", f"h_{x}", 2) for x in ("3pl", "arm", "seg", "anal", "perdas", "descl", "cust", "pag", "prov", "hedge"))
    SENS_POS["tese"] = {}
    for mid, mlab in (("tese_m_med", "Margem líquida do embalador: média de 3 campanhas (€/kg)"), ("tese_m_max", "Margem líquida do embalador: máximo, 2019/20 (€/kg)")):
        put(ws, f"A{r}", mlab, bold=True)
        put(ws, f"B{r}", "=" + preg.ref(mid, 0), fmt=FMT_EKG, bold=True)
        r += 1
        header_row(ws, r, ["Desconto \\ captura da margem do embalador", "=" + preg.ref("tese_cap_1", 0), "=" + preg.ref("tese_cap_2", 0), "=" + preg.ref("tese_cap_3", 0)])
        for col in "BCD":
            ws[f"{col}{r}"].number_format = FMT_PCT
        hr = r
        r += 1
        grid = []
        for dp in ("sens_desc_1", "sens_desc_2", "sens_desc_3"):
            put(ws, f"A{r}", "=" + preg.ref(dp, 0), fmt=FMT_PCT, align=Alignment(horizontal="left"))
            for col in "BCD":
                put(ws, f"{col}{r}", f"={col}${hr}*{preg.ref(mid, 0)}-$A{r}*{ref}-{cup}-({cd})", fmt=FMT_EKG)
            grid.append(r)
            r += 1
        SENS_POS["tese"][mid] = (hr, grid)
        r += 1
    put(ws, f"A{r}", "Todas as células negativas: com 80–90% da margem do embalador não se paga um desconto de 10–15% sobre a referência, "
                     "nem sequer antes do cupão e dos custos da plataforma.", color=MUTED, italic=True)
    r += 2
    # choque de preço no stock
    section_row(ws, r, "Choque de −30% no preço: perda pontual no stock de azeite (2031)", ncols=9)
    r += 1
    header_row(ws, r, ["Cenário", "Stock médio de azeite (€)", "Perda: 100% em aberto (€)", "Perda: 15% em aberto (€)", "Perda 100% ÷ margem anual do azeite"])
    r += 1
    SENS_POS["stock"] = {}
    for s in SCEN:
        put(ws, f"A{r}", SCEN_LONG[SCEN.index(s)])
        put(ws, f"B{r}", "=" + mref(blocks, s, "cg_stock_az", 4), fmt=FMT_EUR)
        put(ws, f"C{r}", f"={preg.ref('sens_preco', 0)}*B{r}", fmt=FMT_EUR)
        put(ws, f"D{r}", f"={preg.ref('exp_aberta', 0)}*C{r}", fmt=FMT_EUR)
        put(ws, f"E{r}", f"=IF({mref(blocks, s, 'mc_az', 4)}>0,-C{r}/{mref(blocks, s, 'mc_az', 4)},\"n/a\")", fmt=FMT_PCT, align=Alignment(horizontal="right"))
        SENS_POS["stock"][s] = r
        r += 1
    r += 1
    section_row(ws, r, "O que tem de ser verdade em 2031 (Base e Recomendado): ver Cenários, secção 3", ncols=9)
    ws.freeze_panes = "B4"
    return ws


# ----------------------------------------------------------------------------- Leia-me e Fontes
def write_leiame(wb, P):
    ws = wb["Sheet"]
    ws.title = "Leia-me"
    title(ws, "Sprout — Modelo financeiro 2027–2031",
          "Gerado por docs/business-plan/modelo/build_model.py a partir de pressupostos.json · data de referência 13-10-2026 · euros correntes, sem IVA")
    for k, v in {"A": 50, "B": 12, "C": 20, "D": 20, "E": 20, "F": 20, "G": 40}.items():
        ws.column_dimensions[k].width = v
    r = 16   # linhas 4–14: tabela de resultados (escrita depois de Cenários)
    blocks = [
        ("Para que serve", [
            "Testar, com números rastreáveis, se a tese do documento-fonte (marketplace B2B de azeite, água e café com desconto de 10–15%, captura de 80–90% "
            "da margem do embalador e investidores a 18 meses) fecha, e qual é o caminho para a viabilidade que a investigação aponta.",
            "Leitores: fundadores, Gonçalo Almeida (jurídico), Pedro Marques (operações, Lagar do Val) e investidores qualificados. Não é uma oferta de investimento.",
        ]),
        ("Resposta curta", [
            "Os três cenários do documento-fonte, com as 53 correções verificadas e sem ajustes, não chegam ao break-even até 2031: a margem por kg não paga a estrutura.",
            "O cenário Recomendado (reestruturado) muda o produto, os canais e os custos como a investigação recomenda. O resultado está na tabela acima e na folha "
            "Cenários; não foi ajustado para dar break-even. A folha Sensibilidade mostra o caso com todas as alavancas no topo dos intervalos da investigação.",
        ]),
        ("Estrutura", [
            "Pressupostos — todos os inputs, por cenário, com fonte; seletor do cenário ativo em C4.",
            "Unit economics — lotes-tipo de ponta a ponta (HoReCa 20 t, distribuidores 30 t, granel 100 t, água 20 paletes) nos quatro cenários; bolo de intermediação; "
            "teste da tese literal; risco de preço; piloto de 2027.",
            "Volumes, Receitas, Pessoal, Custos, DR, Cash flow, Capital & financiamento — o cenário ativo, ano a ano.",
            "Cenários — os quatro cenários lado a lado; o que tem de ser verdade; o que muda no Recomendado; verificações.",
            "Sensibilidade — desconto 10/12,5/15%, volumes ±30%, preço do AOVE −30% e topo dos intervalos (motor completo por caso), para o Base e o Recomendado; "
            "grelha captura × desconto da tese literal; perda no stock com choque de preço.",
            "Fontes — origem de cada pressuposto e registo de verificação. Motor — o cálculo em blocos fixos por cenário e caso (não editar).",
        ]),
        ("Como mudar de cenário", [
            "Pressupostos!C4: escolher Pessimista, Base, Otimista ou Recomendado. Mudam as folhas de detalhe. Cenários, Unit economics e Sensibilidade "
            "mostram sempre cenários fixos.",
            "Para alterar um pressuposto: mudar a célula azul em Pressupostos (o livro recalcula). Para que a alteração passe aos PDF, mudar pressupostos.json e "
            "voltar a correr build_model.py, que regrava resultados.json.",
        ]),
        ("Código de cores", [
            "Azul = input (editar só estas células). Preto = fórmula. Verde = ligação a outra folha. Amarelo = seletor de cenário.",
        ]),
        ("Convenções", [
            "Preços das mercadorias constantes ao nível de out. 2026 (compra back-to-back, neutra ao nível de preço); salários +3%/ano; opex +2%/ano e 5% de contingência.",
            "Capital de giro = nível médio do ano por prazos médios (stock, recebimento, pagamento ao produtor a 30 dias). Obrigações = % do capital de giro ÷ utilização; "
            "0% em 2027 (CSC art. 348.º/2 [L]).",
            "EBITDA inclui o cupão das obrigações como custo direto do lote (a SPV financia stock lote a lote), como no 01 e em calculo_referencia.py; a DR mostra "
            "também o EBITDA sem cupão.",
            "IRC com taxa PME, derramas municipal e estadual e reporte de prejuízos a 65% [L]. Imposto pago no próprio ano (simplificação).",
            "Break-even em mês = cruzamento do EBITDA mensal com zero, interpolando linearmente entre os pontos médios dos anos (estimativa).",
        ]),
        ("Avisos", [
            "Itens [L] (CSC, RGICSF, PRIIPs, IRC, TSU, retenções) estão por confirmar pelo advogado e pelo contabilista certificado: o acesso a fontes primárias "
            "(DRE, EUR-Lex, Portal das Finanças) estava bloqueado na verificação.",
            "Itens NV (referência HoReCa, preço de saída do embalador, comissões, custos de custódia e 3PL) são hipóteses a medir no piloto, não promessas comerciais.",
            "Não existem futuros de azeite (MFAO extinto). Deter stock 18 meses sem cobertura arrisca perdas de 50–70% do lote (carry 19,8% + VaR95 37–57% − margem).",
        ]),
    ]
    for head, paras in blocks:
        put(ws, f"A{r}", head, bold=True, color=ACCENT, align=Alignment(vertical="top"))
        for ptxt in paras:
            put(ws, f"B{r}", ptxt, align=WRAP)
            ws.merge_cells(f"B{r}:G{r}")
            ws.row_dimensions[r].height = 13.5 * max(1, len(ptxt) // 125 + 1)
            r += 1
        r += 1
    return ws, 4


def write_leiame_summary(ws, r):
    section_row(ws, r, "Resultados-chave (ligados à folha Cenários)", ncols=7)
    r += 1
    header_row(ws, r, ["Indicador", "Unidade"] + SCEN_LABEL)
    r += 1
    items = [("EBITDA 2031", "dr_ebitda", 4, "€", FMT_EUR), ("Resultado líquido 2031", "dr_rl", 4, "€", FMT_EUR),
             ("GMV 2031", "gmv_total", 4, "€", FMT_EUR), ("Cash flow acumulado para o acionista, 2027–2031", "cf_acum", 4, "€", FMT_EUR)]
    for lab, lid, y, unit, fmt in items:
        put(ws, f"A{r}", lab)
        put(ws, f"B{r}", unit, color=MUTED)
        for i, s in enumerate(SCEN):
            put(ws, f"{'CDEF'[i]}{r}", f"=Cenários!{YC[y]}{CEN_POS[lid][s]}", fmt=fmt)
        r += 1
    for lab, lid, unit, fmt in (("Necessidade de capital próprio (pico no horizonte)", "k_need", "€", FMT_EUR),
                                ("Break-even do EBITDA", "be_data", "mês", FMT_DATE),
                                ("1.º ano com cash flow para o acionista ≥ 0", "be_ano_cf", "ano", FMT_YEAR)):
        put(ws, f"A{r}", lab, bold=lid == "k_need")
        put(ws, f"B{r}", unit, color=MUTED)
        for i in range(4):
            put(ws, f"{'CDEF'[i]}{r}", f"=Cenários!{'CDEF'[i]}{CEN_POS[lid]}", fmt=fmt, align=Alignment(horizontal="right"), bold=lid == "k_need")
        r += 1
    put(ws, f"A{r}", "Verificações do modelo (Cenários, secção 5)")
    put(ws, f"C{r}", f"=Cenários!C{CEN_POS['estado']}", bold=True, align=Alignment(horizontal="right"))


def write_fontes(wb, P):
    ws = wb.create_sheet("Fontes")
    title(ws, "Fontes", "Origem de cada pressuposto (texto de pressupostos.json → fontes) e registo de verificação. URL completas em docs/research/00-base-factual.md.")
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 150
    r = 4
    section_row(ws, r, "Documentos", ncols=2)
    r += 1
    docs = [
        ("docs/fonte/plano-estrategico-10-13.md", "Documento-fonte (reunião de 13-10): tese, piloto, responsáveis"),
        ("docs/research/00-base-factual.md", "Base factual consolidada, com fontes, etiquetas e registo de verificação (53 correções)"),
        ("docs/research/01-desenho-do-modelo.md", "Desenho do modelo recomendado: aritmética do lote, estrutura jurídica, risco, piloto, cenários"),
        ("docs/research/10-mercado.md", "Mercado, preços e cadeia de valor (azeite, água, café); TAM/SAM/SOM"),
        ("docs/research/20-regulatorio.md", "Regulatório: prospeto, MiFID, PRIIPs, ECSP, RGICSF, CSC, AIFMD"),
        ("docs/research/30-pagamentos-infra.md", "Pagamentos, PSP, KYC/AML, reconciliação, infraestrutura"),
        ("docs/research/40-risco-coberturas.md", "Risco de preço, carry, coberturas, crédito, stress tests, política de cobertura"),
        ("docs/research/50-concorrencia.md", "Concorrência e comparáveis; economia ECSP; fosso"),
        ("docs/business-plan/modelo/pressupostos.json", "Todos os pressupostos (fonte única dos inputs deste livro)"),
        ("docs/business-plan/modelo/calculo_referencia.py", "Cálculo de referência dos três cenários do documento-fonte (este livro reproduz-o)"),
        ("docs/business-plan/modelo/resultados.json", "Valores recalculados deste livro, por cenário e ano (fonte única dos números dos PDF)"),
    ]
    for a, b in docs:
        put(ws, f"A{r}", a)
        put(ws, f"B{r}", b, color=MUTED)
        r += 1
    r += 1
    section_row(ws, r, "Registo de verificação do modelo financeiro", ncols=2)
    r += 1
    for t in P.get("notas_verificacao_modelo_financeiro", []):
        put(ws, f"B{r}", t, align=WRAP)
        ws.row_dimensions[r].height = 15 * max(1, len(t) // 140 + 1)
        r += 1
    r += 1
    put(ws, f"A{r}", "Correção de leitura (Brasil)", bold=True)
    put(ws, f"B{r}", P["cenario_recomendado"]["correcao_brasil"], align=WRAP)
    ws.row_dimensions[r].height = 30
    r += 2
    section_row(ws, r, "Pressupostos: fonte por parâmetro (caminho no JSON)", ncols=2)
    r += 1
    header_row(ws, r, ["Parâmetro", "Fonte / método"])
    r += 1
    for k, v in P["fontes"].items():
        put(ws, f"A{r}", k, align=WRAP)
        put(ws, f"B{r}", v, align=WRAP, color=MUTED)
        ws.row_dimensions[r].height = 15 * max(1, len(v) // 140 + 1)
        r += 1
    return ws


# ----------------------------------------------------------------------------- recálculo
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""


def recalc(path, timeout=300):
    """Recalcula e regrava o .xlsx com o LibreOffice (macro calculateAll + store)."""
    env = os.environ.copy()
    env["SAL_USE_VCLPLUGIN"] = "svp"
    with tempfile.TemporaryDirectory(prefix="sprout-lo-") as prof:
        url = Path(prof).as_uri()
        subprocess.run(["soffice", "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                       capture_output=True, timeout=120, env=env)
        mdir = Path(prof) / "user" / "basic" / "Standard"
        mdir.mkdir(parents=True, exist_ok=True)
        (mdir / "Module1.xba").write_text(MACRO)
        before = os.stat(path).st_mtime_ns
        res = subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation={url}",
                              "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application", str(path)],
                             capture_output=True, text=True, timeout=timeout, env=env)
        if res.returncode != 0 or os.stat(path).st_mtime_ns == before:
            raise RuntimeError("LibreOffice não recalculou o livro: " + (res.stderr or "")[-500:])


ERRS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A", "Err:")


def scan_errors(path):
    wb = load_workbook(path, data_only=True)
    found = []
    nform = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and any(e in c.value for e in ERRS):
                    found.append(f"{ws.title}!{c.coordinate}: {c.value}")
    wbf = load_workbook(path)
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    nform += 1
    return found, nform


# ----------------------------------------------------------------------------- espelho em Python (verificação)
def mirror(P, s, desc=None, vol=0.0, price=0.0, premio_d=None, tp_scen=None):
    """Mesmo motor em Python, para verificar o livro (independente das fórmulas)."""
    az = P["commodities"]["azeite"]
    chs = {"h": az["canais"]["horeca_gourmet"], "d": az["canais"]["distribuidores_export"], "g": az["canais"]["granel_certificado"]}
    cfc = az["custos_fisicos_e_custodia"]
    fin, cu, cx = P["financiamento"], P["custos"], P["capex"]
    ag, cf = P["commodities"]["agua"], P["commodities"]["cafe"]
    ft = P["receitas"]["financiamento_stock_terceiros"]

    def sv(x, y):
        if isinstance(x, dict):
            x = x[s]
        return x[y] if isinstance(x, list) else x

    vols = {"h": az["volume_horeca_gourmet_t"][s], "d": az["volume_distribuidores_export_t"][s], "g": az["volume_granel_certificado_t"][s]}
    out = {k: [] for k in ("ebitda", "cf", "acum", "rl", "gmv", "mc")}
    k_spv = fin["capital_social_spv_eur"]
    prev_eq, prev_res, prev_fte, acum, loss = k_spv, 0.0, 0.0, 0.0, 0.0
    da = [0.0] * 5
    for y in range(5):
        pct, cup, util = fin["pct_capital_giro_azeite_financiado_por_notes"][s][y], fin["cupao_anual"][s][y], fin["utilizacao_media_capital"][s]
        idx = az["preco_indice_granel_eur_kg"][s] * (1 + price)
        gmv_az = m_az = cg_az = 0.0
        r_hd = r_g = 0.0
        for k, ch in chs.items():
            g = idx + sv(ch["premio_lagar_eur_kg"], y)
            if k == "h":
                ref = ch["preco_referencia_eur_kg"][s] * (1 + price)
                dsc = desc if desc is not None else ch["desconto_efetivo"][s]
            elif k == "d":
                ref = ch["preco_referencia_eur_kg"][s] * (1 + price) + (ch["premio_origem_venda_eur_kg"][s] if premio_d is None else premio_d)
                dsc = ch["desconto_efetivo"][s]
            else:
                ref = idx + ch["premio_venda_sobre_indice_eur_kg"]
                dsc = ch["desconto_efetivo"]
            pr = ref * (1 - dsc)
            emb, ds, dr = ch["embalamento_eur_kg"], sv(ch["dias_stock"], y), sv(ch["dias_recebimento"], y)
            dp, lot = cfc["dias_pagamento_produtor"], ch["lote_medio_t"] * 1000
            c = (g + emb + sv(ch["logistica_eur_kg"], y) + ch["armazem_logistico_picking_eur_kg"] + cfc["armazenagem_eur_t_mes"] * ds / 30 / 1000
                 + g * cfc["seguro_stock_pct_ano"] * ds / 365 + ch["analises_por_lote"] * cfc["analise_fisico_quimica_painel_eur"] / lot
                 + cfc["perdas_fisicas_pct_valor"] * g + cfc["desclassificacao_prob_por_trimestre"] * cfc["desclassificacao_penalizacao_eur_kg"] * ds / 90
                 + ch["inspecao_collateral_manager_eur_por_lote"] / lot + ch["pagamentos_bancarios_eur_kg"] + sv(ch["provisao_credito_pct_vendas"], y) * pr
                 + P["custos"]["hedge"]["azeite"]["custo_explicito_pct"] * pr
                 + cup * (g + emb) * (ds + dr - dp) / 365 * pct / util)
            v = vols[k][y] * 1000 * (1 + vol)
            gmv_az += v * pr
            m_az += v * (pr - c)
            cg_az += v * ((g + emb) * (ds - dp) / 365 + pr * dr / 365)
            if k == "g":
                r_g += v * pr
            else:
                r_hd += v * pr
        pal = ag["volume_paletes"][s][y] * (1 + vol)
        gmv_ag = pal * ag["litros_por_palete"] * ag["preco_b2b_referencia_eur_l"] * (1 - ag["desconto_efetivo"])
        m_ag = gmv_ag * ag["comissao_agente_pct_gmv"][s]
        cv = (0 if 2027 + y < cf["ano_inicio"][s] else cf["volume_t"][s][y]) * 1000 * (1 + vol)
        cfin = cup * cf["preco_compra_eur_kg"] * (cf["dias_stock"] + cf["dias_recebimento"] - cf["dias_pagamento"]) / 365 * pct / util
        m_cf = cv * (cf["preco_venda_eur_kg"] - cf["preco_compra_eur_kg"] - cf["logistica_eur_kg"] - cfin - cf["provisao_credito_pct_vendas"] * cf["preco_venda_eur_kg"]
                     - cf["pagamentos_ops_eur_kg"] - cf["custo_cobertura_pct_compras"][s] * cf["preco_compra_eur_kg"])
        cg_cf = cv * (cf["preco_compra_eur_kg"] * cf["dias_stock"] + cf["preco_venda_eur_kg"] * cf["dias_recebimento"] - cf["preco_compra_eur_kg"] * cf["dias_pagamento"]) / 365
        cart = ft["carteira_media_eur"][tp_scen or s][y]
        bruto = cart * (ft["taxa_verificacao_por_utilizacao_pct"] * ft["rotacoes_por_ano"] + ft["taxa_monitorizacao_anual_pct"])
        kgv = cart / (ft["valor_colateral_eur_kg"] * ft["ltv"]) * ft["rotacoes_por_ano"]
        m_tp = bruto - kgv * ft["custo_verificacao_eur_kg_verificado"] * ft["pct_custo_verificacao_suportado_sprout"][s]
        m = m_az + m_ag + m_cf + m_tp + P["receitas"]["dados_e_indice_eur"][s][y]
        pe = cu["pessoal"]
        tot = fte = 0.0
        for fn, v in pe["funcoes"].items():
            n = v["fte"][s][y]
            fte += n
            fac = pe["fator_salarial_fundadores_2027"]["fator_efetivo_2027"] if (y == 0 and fn in FOUNDERS) else 1
            tot += n * v["salario_bruto_anual"] * (1 + pe["atualizacao_salarial_anual"]) ** y * (1 + v.get("variavel_pct", 0)) * fac
        pess = tot * (1 + pe["tsu_entidade_empregadora"] + pe["outros_encargos_pct"]) + fte * pe["subsidio_refeicao_eur_fte_ano"] * (1 + pe["atualizacao_salarial_anual"]) ** y
        ox = cu["opex_nao_salarial_eur"]
        lines = sum(ox[k][s][y] for k in OPEX_KEYS)
        spv = fin["custos_spv"]
        spvc = (spv["setup_eur"] if y == 0 else 0) + spv["anual_eur"] + spv["por_serie_eur"] * spv["series_por_ano"][s][y] + spv["funcoes_independentes_eur"][y]
        sc = cu["seguro_credito"]
        seg = 0.0 if 2027 + y < sc["incluido_no_opex_desde"] else max(sc["premio_minimo_anual_eur"][s][y], sc["premio_pct_volume_seguro"] * r_hd + sc["premio_pct_volume_seguro_granel"] * r_g)
        opex = lines * (1 + ox["contingencia_pct"]) * (1 + ox["inflacao_anual"]) ** y + spvc + seg
        eb = m - pess - opex
        tel = cx["telemetria_eur_por_local"] * cx["locais_novos"][s][y]
        plat = cx["plataforma_desenvolvimento_capitalizavel_eur"][s][y]
        eq = cx["equipamento_informatico_eur_por_nova_fte"] * max(0.0, fte - prev_fte)
        capex = tel + plat + eq
        for amt, life in ((plat, cx["vida_util_anos"]["software"]), (tel, cx["vida_util_anos"]["sensores"]), (eq, cx["vida_util_anos"]["equipamento_informatico"])):
            for j in range(y, min(5, y + life)):
                da[j] += amt / life
        ti = eb - da[y]
        imp = P["impostos"]
        ded = min(loss, imp["reporte_prejuizos"]["limite_deducao_pct_lucro_tributavel"] * max(0, ti))
        loss = loss + max(0, -ti) - ded
        mc_ = max(0, ti) - ded
        irc = imp["irc_taxa_pme_primeiros_50k"] * min(mc_, 50000) + imp["irc_taxa_geral"][str(2027 + y)] * max(0, mc_ - 50000)
        dm = imp["derrama_municipal_max"] * max(0, ti)
        de = sum(t["taxa"] * max(0, min(ti, t["ate"] if t["ate"] else 1e18) - t["de"]) for t in imp["derrama_estadual"])
        tax = irc + dm + de
        cg = cg_az + m_ag * P["capital_de_giro"]["agua"]["dias_recebimento_comissao"] / 365 + cg_cf
        ob = pct * cg / util
        eq_cg = max(cg - pct * cg, k_spv)
        res = fin["reserva_primeira_perda"]["teto_pct_saldo_notes"] * ob
        cfl = eb - capex - tax - (eq_cg - prev_eq) - (res - prev_res)
        acum += cfl
        out["ebitda"].append(eb)
        out["cf"].append(cfl)
        out["acum"].append(acum)
        out["rl"].append(ti - tax)
        out["gmv"].append(gmv_az + gmv_ag + cv * cf["preco_venda_eur_kg"])
        out["mc"].append(m)
        prev_eq, prev_res, prev_fte = eq_cg, res, fte
    pre = sum(P["custos"]["pre_arranque_2026_eur"].values())
    out["need"] = pre + k_spv + P["capital_de_giro"]["caixa_minima_operacional_eur"] + max(0.0, -min(out["acum"]))
    return out


# ----------------------------------------------------------------------------- leitura e JSON
def month_label(m):
    if not isinstance(m, (int, float)):
        return None
    m = int(m)
    y = 2027 + (m - 1) // 12
    return f"{y}-{(m - 1) % 12 + 1:02d}"


def rd(v, nd=0):
    if isinstance(v, (int, float)):
        return round(v, nd) if nd else int(round(v))
    return v


def read_results(path, blocks, lines, P, preg):
    wb = load_workbook(path, data_only=True)
    mo = wb["Motor"]

    def series(key, lid):
        _, row = blocks[key].pos[lid]
        if LINES_BY[lid].scalar:
            return mo[f"C{row}"].value
        return [mo[f"{c}{row}"].value for c in YC]

    def nd_for(unit):
        return {"€": 0, "ano": 0, "data": 0, "t": 1, "FTE": 1, "n.º": 1, "dias": 1, "mês": 1, "meses": 1}.get(unit, 4)

    def rr(x, nd):
        if not isinstance(x, (int, float)):
            return x
        return int(round(x)) if nd == 0 else round(x, nd)

    def all_lines(key):
        out = {}
        for ln_ in lines:
            v = series(key, ln_.lid)
            nd = nd_for(ln_.unit)
            out[ln_.lid] = [rr(x, nd) for x in v] if isinstance(v, list) else rr(v, nd)
        return out

    res = {
        "meta": {
            "gerado_por": "docs/business-plan/modelo/build_model.py",
            "livro": "docs/business-plan/Sprout-Modelo-Financeiro.xlsx",
            "pressupostos": "docs/business-plan/modelo/pressupostos.json",
            "data_referencia": "2026-10-13",
            "unidade": "EUR correntes, sem IVA; volumes em t; preços em EUR/kg",
            "anos": YEARS,
            "cenarios": dict(zip(SCEN, SCEN_LONG)),
            "convencoes": [
                "EBITDA inclui o cupão das obrigações (custo de financiar stock, lote a lote), como no 01 e em calculo_referencia.py; ebitda_sem_cupao à parte",
                "capital de giro = nível médio do ano por prazos médios; obrigações = % do capital de giro / utilização; 0% em 2027 (CSC art. 348.º/2 [L])",
                "necessidade de capital próprio = pré-arranque + capital social da SPV + caixa mínima + vale do cash flow acumulado para o acionista",
                "break-even em mês = cruzamento do EBITDA mensal (run-rate interpolado entre os pontos médios dos anos) com zero; estimativa",
                "valores lidos do livro depois de recalculado pelo LibreOffice (headless); nenhum número foi escrito à mão",
            ],
        },
        "cenarios": {},
    }
    for i, s in enumerate(SCEN):
        A = all_lines(s)
        sc = {
            "nome": SCEN_LONG[i],
            "volumes_t": {"azeite_horeca_gourmet": A["v_h_t"], "azeite_distribuidores_exportacao": A["v_d_t"], "azeite_granel_certificado": A["v_g_t"],
                          "azeite_total": A["v_az_t"], "clientes_horeca_gourmet": A["v_h_cli"], "agua_paletes": A["v_ag_pal"], "cafe_verde": A["v_cf_t"],
                          "quota_producao_pt": A["v_quota"]},
            "gmv_total": A["gmv_total"],
            "gmv_agua": A["gmv_ag"],
            "capital_financiado": {"obrigacoes_em_circulacao": A["ob"], "valor_lotes_comprados": A["lotes_valor"],
                                   "valor_lotes_financiado_por_obrigacoes": A["lotes_fin_obr"], "investidores_estimados": A["inv_n"],
                                   "cupao_pago": A["cup_total"], "reserva_primeira_perda": A["res"]},
            "receitas": {"azeite_horeca_gourmet": A["r_h"], "azeite_distribuidores_exportacao": A["r_d"], "azeite_granel_certificado": A["r_g"],
                         "cafe_verde": A["r_cf"], "comissoes_agua": A["r_ag"], "verificacao_terceiros": A["r_tp"], "dados_indice": A["r_dd"], "total": A["r_total"]},
            "margem_bruta": A["dr_mb"],
            "margem_contribuicao": {"azeite_horeca_gourmet": A["mc_h"], "azeite_distribuidores_exportacao": A["mc_d"], "azeite_granel_certificado": A["mc_g"],
                                    "cafe_verde": A["mc_cf"], "agua": A["mc_ag"], "verificacao_terceiros_liquida": A["mc_tp"], "dados": A["mc_dd"],
                                    "total": A["mc_total"], "azeite_eur_kg": A["mc_az_kg"], "azeite_pct_vendas": A["mc_az_pct"]},
            "margem_por_kg": {"horeca_gourmet": A["h_margem"], "distribuidores_exportacao": A["d_margem"], "granel_certificado": A["g_margem"]},
            "pessoal": A["pess_total"], "fte": A["fte_total"], "opex_nao_salarial": A["ox_total"], "seguro_credito": A["ox_seg"],
            "ebitda": A["dr_ebitda"], "ebitda_sem_cupao": A["dr_ebitda_xc"], "amortizacoes": A["da"], "ebit": A["dr_ebit"], "imposto": A["tx_total"],
            "resultado_liquido": A["dr_rl"], "capex": A["capex"], "capital_giro": A["cg"], "stock_azeite": A["cg_stock_az"],
            "cash_flow_acionista": A["cf_acion"], "cash_flow_acumulado": A["cf_acum"], "rondas_adicionais": A["cx_plug"], "caixa_livre_fim": A["cx_fim"],
            "caixa_total_fim": A["cx_total"], "capital_proprio_levantado_acumulado": A["cx_rondas_acum"],
            "necessidade_capital_proprio": A["k_need"], "ano_minimo_cash_flow_acumulado": A["k_min_ano"], "necessidade_ainda_a_crescer_2031": A["k_ainda"],
            "aplicacao_fundos": {"pre_arranque_e_capital_spv": A["u_pre"], "perdas_operacionais": A["u_ebitda"], "capex": A["u_capex"], "impostos": A["u_imp"],
                                 "capital_giro_reserva_float": A["u_cg"], "caixa_minima": A["u_cmin"], "total": A["u_total"]},
            "pre_seed": A["b_seed"], "gasto_ate_gate_2027_08": A["b_gate"], "gasto_ate_2027_12": A["b_2027"],
            "runway_pre_seed_meses": A["b_runway"],
            "break_even_ebitda_mes": month_label(A["be_mes"]), "break_even_ebitda_mes_indice": A["be_mes"] if isinstance(A["be_mes"], (int, float)) else None,
            "primeiro_ano_ebitda_positivo": A["be_ano_ebitda"], "primeiro_ano_resultado_liquido_positivo": A["be_ano_rl"],
            "primeiro_ano_cash_flow_positivo": A["be_ano_cf"],
            "o_que_tem_de_ser_verdade_2031": {"defice_ebitda": A["wt_gap"][4], "premio_adicional_distribuidores_eur_kg": A["wt_d"][4],
                                              "referencia_horeca_eur_kg": A["wt_h"][4], "volume_granel_t": A["wt_g"][4], "corte_custos_fixos_pct": A["wt_fix"][4],
                                              "carteira_terceiros_eur": A["wt_tp"][4]},
            "linhas": A,
        }
        res["cenarios"][s] = sc
    # unit economics
    ue = wb["Unit economics"]
    uej = {"ano_regime": ue["C3"].value}
    for k, nm in (("h", "horeca_gourmet_20t"), ("d", "distribuidores_exportacao_30t"), ("g", "granel_certificado_100t")):
        rows = UE_POS[k]
        d = {}
        for i, s in enumerate(SCEN):
            ce, ck = "CEGI"[i], "DFHJ"[i]
            dd = {"eur_kg": {}, "eur_lote": {}}
            for key in ("ref", "desc", "price", "gr", "emb", "log", "3pl", "arm", "seg", "anal", "perdas", "descl", "cust", "pag", "prov", "hedge", "cupao", "margem"):
                dd["eur_kg"][key] = rd(ue[f"{ck}{rows[key]}"].value, 4)
                dd["eur_lote"][key] = rd(ue[f"{ce}{rows[key]}"].value)
            dd["margem_pct"] = rd(ue[f"{ck}{rows['margem_pct']}"].value, 4)
            for key in ("cap", "ciclo", "pct", "notes", "sub", "cup", "cupliq", "cupciclo", "retciclo", "res", "livre", "kyc", "noc"):
                v = ue[f"{ce}{rows[key]}"].value
                dd[key] = rd(v, 4) if key in ("pct", "cup", "cupliq", "retciclo") else rd(v)
            d[s] = dd
        uej[nm] = d
    for s_i, s in enumerate(SCEN):
        col = "CDEF"[s_i]
        uej.setdefault("bolo_intermediacao_horeca", {})[s] = {k: rd(ue[f"{col}{UE_POS['pi'][k]}"].value, 4) for k in UE_POS["pi"]}
        uej.setdefault("tese_literal_horeca", {})[s] = {k: (rd(ue[f"{col}{UE_POS['tese'][k]}"].value, 4) if k != "verd" else ue[f"{col}{UE_POS['tese'][k]}"].value)
                                                     for k in UE_POS["tese"]}
        uej.setdefault("risco_preco_horeca", {})[s] = {k: rd(ue[f"{col}{UE_POS['risco'][k]}"].value, 4) for k in UE_POS["risco"]}
        uej.setdefault("piloto_2027_eur_kg", {})[s] = {k: rd(ue[f"{col}{UE_POS['piloto'][k]}"].value, 4) for k in UE_POS["piloto"]}
    agj = {}
    for j, col in enumerate("CDEFGH"):
        name = ["principal_desconto_10", "principal_desconto_15", "agente_pessimista", "agente_base", "agente_otimista", "agente_recomendado"][j]
        agj[name] = {k: (rd(ue[f"{col}{r}"].value, 4) if isinstance(ue[f"{col}{r}"].value, float) and abs(ue[f"{col}{r}"].value) < 50 else rd(ue[f"{col}{r}"].value))
                     for k, r in UE_POS["agua"].items()}
    uej["agua_20_paletes"] = agj
    res["unit_economics"] = uej
    # sensibilidade
    se = wb["Sensibilidade"]
    sens = {"casos": {"desconto_horeca": P["sensibilidades"]["desconto_horeca"], "volume": P["sensibilidades"]["volume"], "preco_aove": P["sensibilidades"]["preco_aove"],
                      "topo_dos_intervalos": P["sensibilidades"]["nota_topo"]}}
    names = {"": "cenario", "__desc_1": "desconto_10", "__desc_2": "desconto_12_5", "__desc_3": "desconto_15",
             "__vol_neg": "volume_menos_30", "__vol_pos": "volume_mais_30", "__preco": "preco_aove_menos_30", "__topo": "topo_dos_intervalos"}
    for s in ("base", "recomendado"):
        d = {}
        for suf, nm in names.items():
            key = s + suf
            d[nm] = {
                "margem_horeca_2029_eur_kg": rd(series(key, "h_margem")[2], 4),
                "ebitda_2029": rd(series(key, "dr_ebitda")[2]),
                "ebitda_2031": rd(series(key, "dr_ebitda")[4]),
                "ebitda": [rd(x) for x in series(key, "dr_ebitda")],
                "cash_flow_acumulado_2031": rd(series(key, "cf_acum")[4]),
                "necessidade_capital_proprio": rd(series(key, "k_need")),
                "break_even_ebitda_mes": month_label(series(key, "be_mes")),
                "margem_media_azeite_2031_eur_kg": rd(series(key, "mc_az_kg")[4], 4),
            }
        sens[s] = d
    tese = {}
    for mid in ("tese_m_med", "tese_m_max"):
        hr, grid = SENS_POS["tese"][mid]
        tese[mid.replace("tese_m_", "margem_embalador_")] = {
            "margem_embalador_eur_kg": PREG.by[mid].vals,
            "captura": [se[f"{c}{hr}"].value for c in "BCD"],
            "desconto": [se[f"A{r}"].value for r in grid],
            "saldo_eur_kg": [[rd(se[f"{c}{r}"].value, 4) for c in "BCD"] for r in grid],
        }
    sens["tese_literal_horeca_base_2029"] = tese
    sens["choque_preco_stock_2031"] = {s: {"stock_azeite": rd(se[f"B{r}"].value), "perda_100_aberto": rd(se[f"C{r}"].value),
                                           "perda_15_aberto": rd(se[f"D{r}"].value), "perda_100_sobre_margem_azeite": rd(se[f"E{r}"].value, 4)}
                                       for s, r in SENS_POS["stock"].items()}
    res["sensibilidade"] = sens
    cen = wb["Cenários"]
    res["verificacao"] = {"estado_verificacoes_internas": cen[f"C{CEN_POS['estado']}"].value}
    return res


# ----------------------------------------------------------------------------- main
PREG = None


def main():
    global PRE_IDS, PREG, LINES_BY
    P = json.load(open(JSON_PATH))
    # coerência: nos cenários do documento-fonte o prémio por canal = prémio global (calculo_referencia.py)
    az = P["commodities"]["azeite"]
    for s in SCEN[:3]:
        for chn in ("horeca_gourmet", "distribuidores_export"):
            assert az["canais"][chn]["premio_lagar_eur_kg"][s] == az["premio_origem_pago_ao_lagar_eur_kg"][s], (s, chn)
    PRE_IDS = ["pre_" + k for k in P["custos"]["pre_arranque_2026_eur"]]
    params = build_params(P)
    PREG = PRegistry(params)
    lines = define_lines()
    LINES_BY = {ln_.lid: ln_ for ln_ in lines}
    assert len(LINES_BY) == len(lines), "ids de linha duplicados"

    wb = Workbook()
    ws_leia, r_leia = write_leiame(wb, P)
    write_pressupostos(wb, params, P)
    det_lay = layout_detail(lines)
    sheets = write_detail(wb, lines, PREG, det_lay)
    write_pessoal_table(sheets["Pessoal"], PREG, P)
    _, blocks = write_motor(wb, lines, PREG)
    write_unit_economics(wb, PREG, blocks)
    write_cenarios(wb, blocks, P, det_lay)
    write_sensibilidade(wb, PREG, blocks)
    write_fontes(wb, P)
    write_leiame_summary(ws_leia, r_leia)
    order = ["Leia-me", "Pressupostos", "Unit economics", "Volumes", "Receitas", "Pessoal", "Custos", "DR", "Cash flow",
             "Capital & financiamento", "Cenários", "Sensibilidade", "Fontes", "Motor"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = ACCENT if ws.title in ("Leia-me", "Cenários") else (BLUE if ws.title == "Pressupostos" else None)
    wb.save(OUT_XLSX)
    print("escrito:", OUT_XLSX)
    if "--no-recalc" in sys.argv:
        return

    t0 = time.time()
    recalc(OUT_XLSX)
    errs, nform = scan_errors(OUT_XLSX)
    print(f"recalculado em {time.time() - t0:.0f}s; fórmulas: {nform}; erros: {len(errs)}")
    if errs:
        for e in errs[:40]:
            print("  ", e)
        sys.exit(1)

    res = read_results(OUT_XLSX, blocks, lines, P, PREG)
    # verificação contra calculo_referencia.py (três cenários do documento-fonte)
    sys.path.insert(0, str(HERE))
    import calculo_referencia as cr
    ver = {"formulas": nform, "erros_de_formula": 0, "tolerancia_relativa": 0.005, "contra_calculo_referencia": {}, "contra_espelho_python": {}}
    worst = 0.0

    def cmp(a, b):
        return abs(a - b) / max(1.0, abs(b))

    for s in SCEN[:3]:
        ref = cr.run(P, s)
        L_ = res["cenarios"][s]["linhas"]
        d = {}
        for ours, theirs in (("dr_ebitda", "ebitda"), ("gmv_total", "gmv_total"), ("mc_total", "margem_contribuicao_total"), ("pess_total", "custos_pessoal"),
                             ("ox_total", "opex_nao_salarial"), ("capex", "capex"), ("cg", "capital_giro_liquido"), ("ob", "obrigacoes_em_circulacao_fim_ano"),
                             ("res", "reserva"), ("cf_acion", "cf"), ("cf_acum", "cash_flow_acionista_acumulado")):
            e = max(cmp(a, b) for a, b in zip(L_[ours], ref[theirs]))
            d[ours] = round(e, 6)
            worst = max(worst, e)
        e = cmp(L_["k_need"], ref["necessidade_capital_proprio"])
        d["k_need"] = round(e, 6)
        worst = max(worst, e)
        ver["contra_calculo_referencia"][s] = d
    ver["contra_calculo_referencia"]["maior_desvio_relativo"] = round(worst, 6)
    worst2 = 0.0
    cases = [(s, {}, s) for s in SCEN]
    for s in ("base", "recomendado"):
        sens = P["sensibilidades"]
        cases += [(s, {"desc": sens["desconto_horeca"][j]}, f"{s}__desc_{j + 1}") for j in range(3)]
        cases += [(s, {"vol": sens["volume"][0]}, f"{s}__vol_neg"), (s, {"vol": sens["volume"][1]}, f"{s}__vol_pos"),
                  (s, {"price": sens["preco_aove"]}, f"{s}__preco"),
                  (s, {"vol": sens["volume"][1], "premio_d": sens["premio_origem_distribuidores_topo"], "tp_scen": "otimista"}, f"{s}__topo")]
    wbv = load_workbook(OUT_XLSX, data_only=True)["Motor"]
    for s, kw, key in cases:
        m = mirror(P, s, **kw)
        got = {lid: [wbv[f"{c}{blocks[key].pos[lid][1]}"].value for c in YC] for lid in ("dr_ebitda", "cf_acum", "dr_rl", "gmv_total", "mc_total")}
        e = max(max(cmp(a, b) for a, b in zip(got["dr_ebitda"], m["ebitda"])), max(cmp(a, b) for a, b in zip(got["cf_acum"], m["acum"])),
                max(cmp(a, b) for a, b in zip(got["dr_rl"], m["rl"])), max(cmp(a, b) for a, b in zip(got["gmv_total"], m["gmv"])),
                max(cmp(a, b) for a, b in zip(got["mc_total"], m["mc"])),
                cmp(wbv[f"C{blocks[key].pos['k_need'][1]}"].value, m["need"]))
        ver["contra_espelho_python"][key] = round(e, 6)
        worst2 = max(worst2, e)
    ver["contra_espelho_python"]["maior_desvio_relativo"] = round(worst2, 6)
    res["verificacao"].update(ver)
    print(f"desvio máx. vs calculo_referencia: {worst:.2e}; vs espelho Python: {worst2:.2e}; verificações internas: {res['verificacao']['estado_verificacoes_internas']}")
    if worst > 0.005 or worst2 > 0.005 or res["verificacao"]["estado_verificacoes_internas"] != "OK":
        json.dump(res, open(OUT_JSON, "w"), indent=1, ensure_ascii=False)
        sys.exit("verificação falhou")
    json.dump(res, open(OUT_JSON, "w"), indent=1, ensure_ascii=False)
    print("escrito:", OUT_JSON)


if __name__ == "__main__":
    main()
